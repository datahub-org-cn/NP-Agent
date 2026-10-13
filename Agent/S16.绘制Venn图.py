# -*- coding: utf-8 -*-
"""
S16.绘制Venn图.py
=================
绘制"去重后的药物作用 Gene symbol"与"疾病 Gene symbol"的 Venn 图。

数据源:
    药物靶点: S12.全部GeneSymbol汇总.xlsx (GeneSymbol汇总, 来自S8/S9/S11, 563个)
    疾病靶点: S14.全部疾病GeneSymbol-纳入.xlsx (GeneSymbol汇总, 来自OpenTargets+GeneCards, 2700个)

输出:
    S16.Venn图-药物vs疾病.png   —— Venn 图(含各集合与交集数量)
    S16.Venn图-药物vs疾病.xlsx  —— 交集基因(核心靶点) / 统计

用法: python S16.绘制Venn图.py
"""
import os
import re
import sys

S12 = "S12.全部GeneSymbol汇总.xlsx"
S14 = "S14.全部疾病GeneSymbol-纳入.xlsx"
PNG_OUT = "S16.Venn图-药物vs疾病.png"
XLSX_OUT = "S16.Venn图-药物vs疾病.xlsx"


def norm_gene(v):
    s = str(v or "").strip()
    return s.upper() if s and s.upper() != "NA" else ""


def read_genes(path, sheet, src_col=None, src_val=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = {}
    for r in rows[1:]:
        d = {header[i]: ("" if v is None else v) for i, v in enumerate(r)}
        g = norm_gene(d.get("Gene symbol"))
        if not g:
            continue
        out[g] = d
    wb.close()
    return out


def setup_chinese_font():
    from matplotlib import font_manager
    import matplotlib.pyplot as plt
    available = {f.name for f in font_manager.fontManager.ttflist}
    for name in ("Microsoft YaHei", "SimHei", "SimSun", "Arial Unicode MS",
                 "PingFang SC", "Noto Sans CJK SC"):
        if name in available:
            plt.rcParams["font.sans-serif"] = [name] + plt.rcParams["font.sans-serif"]
            break
    plt.rcParams["axes.unicode_minus"] = False
    return plt.rcParams["font.sans-serif"][0]


def main():
    if not all(os.path.exists(p) for p in (S12, S14)):
        print("错误: 缺少输入文件")
        for p in (S12, S14):
            if not os.path.exists(p):
                print("  找不到:", p)
        sys.exit(1)

    drug_map = read_genes(S12, "GeneSymbol汇总")
    dis_map = read_genes(S14, "GeneSymbol汇总")
    drug = set(drug_map)
    dis = set(dis_map)
    inter = drug & dis
    only_drug = drug - dis
    only_dis = dis - drug
    print(f"药物靶点基因: {len(drug)}")
    print(f"疾病靶点基因: {len(dis)}")
    print(f"交集(核心靶点): {len(inter)}")
    print(f"仅药物: {len(only_drug)} / 仅疾病: {len(only_dis)}")

    font_used = setup_chinese_font()
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib_venn import venn2

    fig, ax = plt.subplots(figsize=(8, 6), dpi=100)
    v = venn2(subsets=(len(only_drug), len(only_dis), len(inter)),
              set_labels=("药物靶点\n(S8/S9/S11)", "疾病靶点\n(OpenTargets+GeneCards)"),
              set_colors=("#4C72B0", "#DD8452"), alpha=0.75, ax=ax)
    ax.set_title(f"药物靶点 vs 疾病靶点 Venn图\n交集核心靶点 = {len(inter)} 个", fontsize=13)
    plt.tight_layout()
    fig.savefig(PNG_OUT, dpi=300, bbox_inches="tight")
    print(f"已写出 {PNG_OUT} (字体: {font_used})")
    plt.close(fig)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # 交集基因(核心靶点)
    ws0 = wb.active
    ws0.title = "交集基因(核心靶点)"
    ws0.append(["Gene symbol", "药物侧来源", "疾病侧来源"])
    for g in sorted(inter):
        dd = drug_map.get(g, {})
        dm = dis_map.get(g, {})
        ws0.append([g, dd.get("来源", ""), dm.get("来源", "")])
    style_header(ws0, 3)
    for i, c in enumerate(["Gene symbol", "药物侧来源", "疾病侧来源"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 16
    ws0.column_dimensions["A"].width = 18

    # 统计
    ws1 = wb.create_sheet("统计")
    ws1.append(["统计项", "数量"])
    for k, v in [("药物靶点基因数", len(drug)), ("疾病靶点基因数", len(dis)),
                 ("交集(核心靶点)", len(inter)), ("仅药物", len(only_drug)),
                 ("仅疾病", len(only_dis))]:
        ws1.append([k, v])
    style_header(ws1, 2)
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 10

    wb.save(XLSX_OUT)
    print(f"已写出 {XLSX_OUT}")


if __name__ == "__main__":
    main()
