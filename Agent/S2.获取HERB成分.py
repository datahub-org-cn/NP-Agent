# -*- coding: utf-8 -*-
"""
获取HERB成分.py  (增强版)
=========================
从 HERB 2.0 (http://47.92.70.12/) 获取 S1复方表.xlsx 中各药物的成分(ingredients)，
逐个校验(存在性/名称/分子式)，并补充 PubChem CID、PubChem 名、CAS、InChI/InChIKey、
IUPAC、异构SMILES、平均分子量、LogP、DL、OB、TCMSP MOL ID 等字段。
结果另存为单独文件。

输入文件:
    S1复方表.xlsx   —— 列1中文标准名、列2中文别名(顿号分隔)、列3拉丁名、列4英文名

输出文件:
    S2.HERB成分表.xlsx —— 药物汇总(含匹配名称) + 成分校验 + 每药一个工作表

HERB 2.0 接口:
    POST http://47.92.70.12/chedi/api/   body: {"func_name": ...}
        search_api  : keyword=名称, label=Herb/Ingredient -> 搜索
        detail_api  : key_id=HERBxxxxxx, label=Herb      -> 草药详情(herb_ingredient)
                      key_id=HBINxxxxxx, label=Ingredient -> 成分详情(summary含PubChem等)

搜索规则:
    优先按「中文标准名」搜索; 未命中则依次按「中文别名」搜索(顿号分隔)。

替代规则:
    赶黄草(扯根菜, Penthorum chinense) HERB 未收录；S1表中其别名为「珍珠菜、扯根菜」,
    「珍珠菜」(Lysimachia clethroides, HERB006968) 由别名搜索自动命中。
    若个别药物标准名+别名均无法命中, 可在 SUBSTITUTE 中硬编码指定 HERB ID。

用法:
    python 获取HERB成分.py
"""
import json
import os
import re
import sys
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed

API = "http://47.92.70.12/chedi/api/"
S1_XLSX = "S1复方表.xlsx"
XLSX_OUT = "S2.HERB成分表.xlsx"

# 硬编码替代规则(可扩展): 若某药 HERB 未收录、且标准名+别名均无法命中时, 可在此手动指定 HERB ID。
# 当前 S1 中「赶黄草」别名为「珍珠菜、扯根菜」, HERB 中「珍珠菜」= HERB006968, 由别名搜索自动命中, 无需硬编码。
SUBSTITUTE = {}


def call_api(func_name, **kwargs):
    data = json.dumps({"func_name": func_name, **kwargs}).encode("utf-8")
    req = urllib.request.Request(API, data=data,
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return json.loads(resp.read().decode("utf-8"))


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, dict):
        return str(v.get("title", ""))
    return str(v)


def search_herb(cn_name):
    res = call_api("search_api", keyword=cn_name, label="Herb")
    rows = res.get("res_data", [])
    for row in rows[1:]:
        if len(row) > 2 and to_str(row[2]) == cn_name:
            return row, True
    if len(rows) > 1:
        return rows[1], False
    return None, False


def find_herb(names):
    """依次用多个名称搜索 HERB 草药, 返回 (herb行, 是否精确匹配, 命中的名称)"""
    for name in names:
        herb, exact = search_herb(name)
        if herb:
            return herb, exact, name
    return None, False, ""


def get_herb_detail(hid):
    return call_api("detail_api", key_id=hid, label="Herb")


def process_ingredient(iid, name, formula):
    """校验 + 详情: 返回(校验字段, 详情字段)"""
    # 校验: 成分ID回查
    res = call_api("search_api", keyword=iid, label="Ingredient")
    exists, herb_name, herb_formula = False, "", ""
    for row in res.get("res_data", [])[1:]:
        if to_str(row[0]) == iid:
            exists = True
            herb_name = to_str(row[1]) if len(row) > 1 else ""
            herb_formula = to_str(row[3]) if len(row) > 3 else ""
            break
    name_ok = exists and (herb_name == name)
    formula_ok = exists and (herb_formula == formula)
    if exists and name_ok and formula_ok:
        status = "通过"
    elif not exists:
        status = "HERB中未检索到该成分ID"
    else:
        status = "不一致"

    # 详情: 补 PubChem/CAS/InChI 等
    d = {}
    try:
        det = call_api("detail_api", key_id=iid, label="Ingredient")
        summary = det.get("summary", [[], []])
        if summary and len(summary) >= 2:
            d = dict(zip(summary[0], summary[1]))
    except Exception:
        d = {}
    return {
        "exists": exists, "herb_name": herb_name, "herb_formula": herb_formula,
        "name_ok": name_ok, "formula_ok": formula_ok, "status": status,
        "detail": d,
    }


def read_drugs_from_s1(path):
    """读取 S1复方表.xlsx: 中文标准名 / 中文别名(顿号分隔) / 拉丁名 / 英文名"""
    from openpyxl import load_workbook
    from collections import Counter

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    drugs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        cn = str(row[0]).strip()
        if not cn:
            continue
        alias_raw = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        latin = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        en = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        aliases = [a for a in re.split(r"[、，,;；/／\s]+", alias_raw) if a]
        drugs.append({"标准名": cn, "别名": aliases, "拉丁名": latin, "英文名": en})
    cnt = Counter(d["标准名"] for d in drugs)
    for d in drugs:
        d["出现次数"] = cnt[d["标准名"]]
    return drugs


def main():
    if not os.path.exists(S1_XLSX):
        print(f"错误: 找不到 {S1_XLSX}")
        sys.exit(1)

    drugs = read_drugs_from_s1(S1_XLSX)
    print(f"读取到 {len(drugs)} 味药物: " + "、".join(d["标准名"] for d in drugs))

    summary_rows, ingredient_rows, all_ing = [], [], []
    cols = ["药物", "出现次数", "HERB ID", "匹配名称", "成分ID", "成分名", "别名",
            "PubChem CID", "PubChem 名称", "CAS", "分子式", "Canonical SMILES",
            "异构SMILES", "InChI", "InChIKey", "IUPAC名称", "MW", "Hacc", "Hdon",
            "LogP", "RBN", "DL", "OB(%)", "TCMSP MOL", "SymMap", "DrugBank",
            "NPASS", "HIT 2.0", "TCMID", "TCM-ID", "Morgan指纹", "MolWt_judge",
            "Hacc_judge", "Hdon_judge", "LogP_judge", "RBN_judge", "URL",
            "HERB中是否存在", "HERB标准名", "名称一致", "HERB分子式", "分子式一致",
            "校验结果", "替代说明"]

    for idx, d in enumerate(drugs, 1):
        cn = d["标准名"]
        aliases = d["别名"]
        latin = d["拉丁名"]
        times = d["出现次数"]
        print(f"[{idx}/{len(drugs)}] 处理 {cn} ...", flush=True)
        sub_hid, sub_note = SUBSTITUTE.get(cn, (None, ""))

        try:
            if sub_hid:
                herb, exact, hid, matched_name = None, False, sub_hid, cn
                note = sub_note
            else:
                herb, exact, matched_name = find_herb([cn] + aliases)
                if not herb:
                    summary_rows.append({"药物": cn, "出现次数": times, "HERB ID": "",
                                         "中文名匹配": "否(HERB未收录)", "匹配名称": "",
                                         "拼音": "", "中文名": "", "英文名": "", "拉丁名": "",
                                         "成分数": 0, "校验结果": "HERB 中未找到", "替代说明": ""})
                    print(f"   {cn}: HERB 中未找到 (标准名及别名均未命中)", flush=True)
                    continue
                hid = to_str(herb[0])
                matched_name = matched_name or cn
                note = ""
                print(f"   {cn}: 按「{matched_name}」命中 {hid}", flush=True)

            obj = get_herb_detail(hid)
            summary = obj.get("summary", [[], []])
            meta = dict(zip(summary[0], summary[1])) if summary and len(summary) >= 2 else {}
            ing = obj.get("herb_ingredient", [])
            header = ing[0] if ing else []
            items = [dict(zip(header, r)) for r in ing[1:]]

            # 并行: 校验 + 详情
            procs = {}
            with ThreadPoolExecutor(max_workers=8) as ex:
                futs = {}
                for r in items:
                    iid = to_str(r.get("Ingredient id") or r.get("Ingredient ID"))
                    fut = ex.submit(process_ingredient, iid,
                                    to_str(r.get("Ingredient name")),
                                    to_str(r.get("Molecular formula")))
                    futs[fut] = iid
                done_n = 0
                for fut in as_completed(futs):
                    procs[futs[fut]] = fut.result()
                    done_n += 1
                    if done_n % 100 == 0:
                        print(f"   进度 {done_n}/{len(items)}", flush=True)

            verified = []
            for r in items:
                iid = to_str(r.get("Ingredient id") or r.get("Ingredient ID"))
                p = procs.get(iid, {})
                det = p.get("detail", {})
                verified.append({
                    "药物": cn, "出现次数": times, "HERB ID": hid, "匹配名称": matched_name,
                    "成分ID": iid,
                    "成分名": to_str(r.get("Ingredient name")),
                    "别名": to_str(r.get("Ingredient alias name")),
                    "PubChem CID": to_str(det.get("PubChem")),
                    "PubChem 名称": to_str(det.get("PubChem name")),
                    "CAS": to_str(det.get("CAS")),
                    "分子式": to_str(r.get("Molecular formula")),
                    "Canonical SMILES": to_str(r.get("Canonical smiles")),
                    "异构SMILES": to_str(det.get("Isomeric smiles")),
                    "InChI": to_str(det.get("InChI")),
                    "InChIKey": to_str(det.get("InChIKey")),
                    "IUPAC名称": to_str(det.get("IUPAC Name")),
                    "MW": to_str(det.get("The average molecular weight")),
                    "Hacc": to_str(det.get("Number of Hydrogen Bond Acceptors")),
                    "Hdon": to_str(det.get("Number of Hydrogen Bond Donors ")),
                    "LogP": to_str(det.get("Wildman-Crippen LogP value")),
                    "RBN": to_str(det.get("Number of Rotatable Bonds")),
                    "DL": to_str(det.get("Drug-likeness")),
                    "OB(%)": to_str(det.get("OB score")),
                    "TCMSP MOL": to_str(det.get("TCMSP")),
                    "SymMap": to_str(det.get("SymMap")),
                    "DrugBank": to_str(det.get("DrugBank")),
                    "NPASS": to_str(det.get("NPASS")),
                    "HIT 2.0": to_str(det.get("HIT 2.0")),
                    "TCMID": to_str(det.get("TCMID")),
                    "TCM-ID": to_str(det.get("TCM-ID")),
                    "Morgan指纹": to_str(det.get("Morgan fingerprint")),
                    "MolWt_judge": to_str(det.get("MolWt_judge")),
                    "Hacc_judge": to_str(det.get("NumHAcceptors_judge")),
                    "Hdon_judge": to_str(det.get("NumHDonors_judge")),
                    "LogP_judge": to_str(det.get("MolLogP_judge")),
                    "RBN_judge": to_str(det.get("NumRotatableBonds_judge")),
                    "URL": to_str(det.get("URL")),
                    "HERB中是否存在": "是" if p.get("exists") else "否",
                    "HERB标准名": p.get("herb_name", ""),
                    "名称一致": "是" if p.get("name_ok") else "否",
                    "HERB分子式": p.get("herb_formula", ""),
                    "分子式一致": "是" if p.get("formula_ok") else "否",
                    "校验结果": p.get("status", ""),
                    "替代说明": note,
                })
            ingredient_rows.extend(verified)
            all_ing.append((cn, hid, verified))

            if sub_hid:
                match_txt = "替代(珍珠菜)"
            elif matched_name == cn:
                match_txt = "是"
            else:
                match_txt = f"是(别名:{matched_name})"
            summary_rows.append({
                "药物": cn, "出现次数": times, "HERB ID": hid,
                "中文名匹配": match_txt, "匹配名称": matched_name,
                "拼音": to_str(meta.get("Herb pinyin name")),
                "中文名": to_str(meta.get("Herb Chinese name")),
                "英文名": to_str(meta.get("Herb English name")),
                "拉丁名": to_str(meta.get("Herb Latin name")),
                "成分数": len(items),
                "校验结果": f"通过{'(替代)' if sub_hid else ''}",
                "替代说明": note,
            })
            print(f"   {cn} -> {hid}: {len(items)} 个成分 (匹配: {matched_name})", flush=True)
        except Exception as ex:
            summary_rows.append({"药物": cn, "出现次数": times, "HERB ID": "",
                                 "中文名匹配": "", "匹配名称": "", "拼音": "", "中文名": "",
                                 "英文名": "", "拉丁名": "", "成分数": 0,
                                 "校验结果": f"出错: {ex}", "替代说明": ""})
            print(f"   {cn} 出错: {ex}", flush=True)


    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("未安装 openpyxl，跳过 xlsx 输出")
        return

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "药物汇总"
    sum_cols = ["药物", "出现次数", "HERB ID", "中文名匹配", "匹配名称", "拼音", "中文名",
                "英文名", "拉丁名", "成分数", "校验结果", "替代说明"]
    ws0.append(sum_cols)
    for s in summary_rows:
        ws0.append([s.get(c, "") for c in sum_cols])
    style_header(ws0, len(sum_cols))
    for i, c in enumerate(sum_cols, 1):
        ws0.column_dimensions[get_column_letter(i)].width = 16

    ws1 = wb.create_sheet("成分校验")
    ws1.append(cols)
    for r in ingredient_rows:
        ws1.append([r.get(c, "") for c in cols])
    style_header(ws1, len(cols))
    for i, c in enumerate(cols, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 16

    for cn, hid, verified in all_ing:
        sheet_title = "赶黄草(珍珠菜)" if cn == "赶黄草" else cn[:31]
        ws = wb.create_sheet(title=sheet_title)
        ws.append(["成分ID", "成分名", "PubChem CID", "CAS", "InChIKey", "分子式",
                   "Canonical SMILES", "IUPAC名称", "MW", "LogP", "DL", "OB(%)", "TCMSP MOL", "校验结果"])
        for v in verified:
            ws.append([v["成分ID"], v["成分名"], v["PubChem CID"], v["CAS"],
                       v["InChIKey"], v["分子式"], v["Canonical SMILES"],
                       v["IUPAC名称"], v["MW"], v["LogP"], v["DL"], v["OB(%)"], v["TCMSP MOL"], v["校验结果"]])
        style_header(ws, 14)
        for i, w in enumerate([14, 40, 12, 26, 34, 14, 50, 40, 14, 10, 12, 10, 12, 12], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_OUT)
    print(f"已写出 {XLSX_OUT}")


if __name__ == "__main__":
    main()
