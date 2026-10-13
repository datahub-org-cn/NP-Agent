# -*- coding: utf-8 -*-
"""
S6.获取药物级靶点_TCMSP.py
==========================
从 TCMSP 数据库获取复方中每味药的【药物级靶点】。
以 S4.TCMSP成分表-纳入.xlsx 中该药物的最终纳入成分为准，
取每个 MOL 的靶点(molecule.php 内嵌 var tar)并合并去重为该药物的药物级靶点。

输入: S4.TCMSP成分表-纳入.xlsx (工作表: 纳入成分)
输出: S6.药物级靶点_TCMSP.xlsx (工作表: 药物汇总 / 靶点明细)

用法: python S6.获取药物级靶点_TCMSP.py
"""
import json
import os
import re
import sys
import time
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

BASE = "https://www.tcmsp-e.com/"
MOLECULE_URL = BASE + "molecule.php"
INPUT = "S4.TCMSP成分表-纳入.xlsx"
SHEET = "纳入成分"
OUT = "S6.药物级靶点_TCMSP.xlsx"
MAX_WORKERS = 6

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
}

DETAIL_COLS = ["药物", "靶点ID", "靶点名", "DrugBank ID", "涉及分子数", "涉及分子", "来源"]


def http_get(url, timeout=60, retries=3):
    last_err = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                raw = resp.read()
            for enc in ("utf-8", "gb18030", "latin-1"):
                try:
                    return raw.decode(enc)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", errors="replace")
        except Exception as ex:
            last_err = ex
            time.sleep(1.5 * (attempt + 1))
    raise last_err


def extract_json_array(html, key):
    m = re.search(re.escape(key) + r"\s*[=:]?\s*(\[)", html)
    if not m:
        return None
    start = m.start(1)
    depth, in_str, esc = 0, False, False
    i = start
    while i < len(html):
        ch = html[i]
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch == "[":
                depth += 1
            elif ch == "]":
                depth -= 1
                if depth == 0:
                    break
        i += 1
    try:
        return json.loads(html[start:i + 1])
    except Exception:
        return None


def get_molecule_targets(mol_id):
    """按 MOL_ID 获取分子靶点列表, 每个元素含 target_ID/target_name/drugbank_ID/SVM_score/RF_score"""
    num = int(mol_id[3:]) if str(mol_id).startswith("MOL") else str(mol_id)
    html = http_get(f"{MOLECULE_URL}?qn={num}")
    tar = extract_json_array(html, "var tar")
    return tar or []


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def write_xlsx(summary_rows, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "药物汇总"
    ws0.append(["药物", "成分数", "药物级靶点数"])
    for r in summary_rows:
        ws0.append([r["药物"], r["成分数"], r["靶点数"]])
    ws0.append(["合计", sum(r["成分数"] for r in summary_rows),
                sum(r["靶点数"] for r in summary_rows)])
    style_header(ws0, 3)
    for i, c in enumerate(["药物", "成分数", "药物级靶点数"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14

    ws1 = wb.create_sheet("靶点明细")
    ws1.append(DETAIL_COLS)
    for r in detail_rows:
        ws1.append([r.get(c, "") for c in DETAIL_COLS])
    style_header(ws1, len(DETAIL_COLS))
    for i, c in enumerate(DETAIL_COLS, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 12
    ws1.column_dimensions[get_column_letter(DETAIL_COLS.index("靶点名") + 1)].width = 34
    ws1.column_dimensions[get_column_letter(DETAIL_COLS.index("涉及分子") + 1)].width = 40
    ws1.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 靶点明细 {len(detail_rows)} 行")


def main():
    if not os.path.exists(INPUT):
        print(f"错误: 找不到 {INPUT}")
        sys.exit(1)
    rows = read_xlsx(INPUT, SHEET)
    print(f"读取 {INPUT} 工作表[{SHEET}]: {len(rows)} 行")

    # 每药 -> MOL_ID 列表
    drug_mols = defaultdict(list)
    for r in rows:
        mol = str(r.get("MOL_ID", "")).strip()
        drug = str(r.get("药物", "")).strip()
        if mol and drug:
            if mol not in drug_mols[drug]:
                drug_mols[drug].append(mol)

    # 全局唯一 MOL_ID 抓取靶点
    all_mols = sorted({m for lst in drug_mols.values() for m in lst})
    print(f"唯一 MOL_ID: {len(all_mols)} 个, 抓取分子靶点 ...", flush=True)
    mol_tar = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = {ex.submit(get_molecule_targets, m): m for m in all_mols}
        done_n = 0
        for fut in as_completed(futs):
            mol = futs[fut]
            try:
                mol_tar[mol] = fut.result()
            except Exception as e:
                mol_tar[mol] = []
                print(f"   {mol} 靶点抓取失败: {e}", flush=True)
            done_n += 1
            if done_n % 20 == 0:
                print(f"   进度 {done_n}/{len(all_mols)}", flush=True)

    summary_rows, detail_rows = [], []
    for drug, mols in drug_mols.items():
        agg = defaultdict(list)   # 靶点ID -> [MOL_ID]
        for mol in mols:
            for t in mol_tar.get(mol, []):
                tid = str(t.get("target_ID", "")).strip()
                if tid:
                    agg[tid].append(mol)
        for tid, ms in agg.items():
            t0 = None
            for mol in ms:
                for t in mol_tar.get(mol, []):
                    if str(t.get("target_ID", "")).strip() == tid:
                        t0 = t
                        break
                if t0:
                    break
            detail_rows.append({
                "药物": drug, "靶点ID": tid,
                "靶点名": str(t0.get("target_name", "")),
                "DrugBank ID": str(t0.get("drugbank_ID", "")),
                "涉及分子数": len(ms), "涉及分子": "; ".join(ms),
                "来源": "TCMSP药物级靶点",
            })
        summary_rows.append({"药物": drug, "成分数": len(mols), "靶点数": len(agg)})
        print(f"   {drug}: {len(mols)} 个成分 -> {len(agg)} 个药物级靶点", flush=True)

    write_xlsx(summary_rows, detail_rows)
    print("=== S6 TCMSP 药物级靶点完成 ===")


if __name__ == "__main__":
    main()
