# -*- coding: utf-8 -*-
"""
S19.核心基因筛选.py
===================
从 STRING PPI 网络（S18 输出）中筛选"核心基因"。

最终方案（已确认，详见 S19方案.html）:
    核心基因 = 方法C(多指标共识≥3) ∩ 综合评分CS排名前10%   (即 D = A∩C, 约36个)
    并把 方法A(Top-N) / 方法B(统计阈值) / 方法C(共识) 作为敏感性分析一并输出。

中心性指标(原理与公式见 S19方案.html):
    度 DC = k/(N-1);  介数 BC = Σ σ_st(i)/σ_st;  接近 CC = (N-1)/Σ d(i,j);
    特征向量 EC (A·x=λ_max·x);  k-核 KC (k-core)
综合评分: CS = 0.30·DĈ + 0.20·BĈ + 0.20·KĈ + 0.15·CĈ + 0.15·EĈ (各指标Min-Max归一化)

输入: S18.STRING-PPI-score700.xlsx (工作表 PPI边)  [可用 --input 指定其它版本]
输出:
    S19.核心基因.xlsx   (核心基因评分 / 方法对比 / 参数)
    S19.核心基因列表.txt(最终核心基因)
    S19.核心基因子网络.png (Cytoscape 渲染, 若可用)

用法:
    python S19.核心基因筛选.py
    python S19.核心基因筛选.py --input S18.STRING-PPI-score900.xlsx --alpha-top 0.15
"""
import argparse
import json
import math
import os
import statistics
import sys
import urllib.request
from collections import Counter

import networkx as nx

DEFAULT_INPUT = "S18.STRING-PPI-score700.xlsx"
XLSX_OUT = "S19.核心基因.xlsx"
TXT_OUT = "S19.核心基因列表.txt"
PNG_OUT = "S19.核心基因子网络.png"
CYTO = "http://127.0.0.1:1234/v1"

WEIGHTS = {"DC": 0.30, "BC": 0.20, "CC": 0.15, "EC": 0.15, "KC": 0.20}


def read_edges(path, sheet="PPI边"):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rows[0]]
    ia, ib = hdr.index("节点A"), hdr.index("节点B")
    isc = hdr.index("combined_score") if "combined_score" in hdr else None
    edges = []
    for r in rows[1:]:
        a, b = str(r[ia]).strip(), str(r[ib]).strip()
        if not a or not b:
            continue
        try:
            sc = float(r[isc]) if isc is not None else 0.4
        except (TypeError, ValueError):
            sc = 0.4
        edges.append((a, b, sc))
    wb.close()
    return edges


def norm(d):
    mn, mx = min(d.values()), max(d.values())
    if mx == mn:
        return {k: 1.0 for k in d}
    return {k: (v - mn) / (mx - mn) for k, v in d.items()}


def rest(method, path, payload=None, timeout=180):
    url = CYTO + path
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(url, data=data, method=method,
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def push_cytoscape(core_genes, edges, out_png):
    """核心基因诱导子网络推入 Cytoscape 并导出 PNG"""
    core_set = set(core_genes)
    sub_edges = [(a, b, s) for a, b, s in edges if a in core_set and b in core_set]
    node_objs = [{"data": {"id": n, "name": n, "label": n}} for n in core_genes]
    edge_objs = [{"data": {"source": a, "target": b, "score": s, "interaction": "PPI"}}
                 for a, b, s in sub_edges]
    net = json.loads(rest("POST", "/networks", {
        "data": {"name": "S19-核心基因子网络"},
        "elements": {"nodes": node_objs, "edges": edge_objs},
    }).decode())
    suid = net["networkSUID"]
    style = {
        "title": "S19_CoreStyle",
        "defaults": [
            {"visualProperty": "NODE_FILL_COLOR", "value": "#C00000"},
            {"visualProperty": "NODE_LABEL_COLOR", "value": "#000000"},
            {"visualProperty": "NODE_SIZE", "value": 40},
            {"visualProperty": "EDGE_STROKE_UNSELECTED_PAINT", "value": "#999999"},
            {"visualProperty": "NODE_LABEL_FONT_SIZE", "value": 12},
        ],
        "mappings": [
            {"mappingType": "passthrough", "mappingColumn": "label", "mappingColumnType": "String",
             "visualProperty": "NODE_LABEL"},
        ],
    }
    try:
        title = json.loads(rest("POST", "/styles", style).decode()).get("title", "S19_CoreStyle")
        rest("POST", "/commands/vizmap/apply", {"network": f"SUID:{suid}", "styles": title})
    except Exception:
        pass
    try:
        rest("POST", "/commands/layout/cose", {"network": f"SUID:{suid}"})
        views = json.loads(rest("GET", f"/networks/{suid}/views").decode())
        if views:
            rest("POST", "/commands/view/fit%20content", {"view": f"SUID:{views[0]}"})
            abs_out = os.path.abspath(out_png).replace("\\", "/")
            rest("POST", "/commands/view/export%20png",
                 {"outputFile": abs_out, "view": f"SUID:{views[0]}"})
            print(f"已导出 {out_png} (核心基因子网络, Cytoscape)")
    except Exception as e:
        print("Cytoscape 渲染失败(跳过):", e)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default=DEFAULT_INPUT, help="PPI 输入 xlsx")
    ap.add_argument("--alpha-top", type=float, default=0.10, help="方法A/最终: CS Top-N 比例(默认0.10)")
    ap.add_argument("--k", type=float, default=1.0, help="方法B: 标准差倍数(默认1.0)")
    ap.add_argument("--alpha-cons", type=float, default=0.20, help="方法C: 单指标候选比例(默认0.20)")
    ap.add_argument("--cons-thr", type=int, default=3, help="方法C: 共识阈值(默认3)")
    ap.add_argument("--no-cytoscape", action="store_true", help="跳过 Cytoscape 子网络")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        print(f"错误: 找不到 {args.input}")
        sys.exit(1)
    edges = read_edges(args.input)
    print(f"读取 {args.input}: {len(edges)} 条边")

    G = nx.Graph()
    for a, b, s in edges:
        G.add_edge(a, b, weight=s)
    lcc = max(nx.connected_components(G), key=len)
    G = G.subgraph(lcc).copy()
    N = G.number_of_nodes()
    print(f"网络: {N} 节点, {G.number_of_edges()} 边 (LCC)")

    # ---- 中心性 ----
    dc = nx.degree_centrality(G)
    bc = nx.betweenness_centrality(G)
    cc = nx.closeness_centrality(G)
    ec = nx.eigenvector_centrality(G, max_iter=1000)
    kc = nx.core_number(G)
    metrics = {"DC": dc, "BC": bc, "CC": cc, "EC": ec, "KC": kc}
    normed = {m: norm(d) for m, d in metrics.items()}

    cs = {n: sum(WEIGHTS[m] * normed[m][n] for m in WEIGHTS) for n in G.nodes()}
    ranked = sorted(cs, key=cs.get, reverse=True)
    rank = {n: i + 1 for i, n in enumerate(ranked)}

    # ---- 方法A ----
    n_a = max(15, math.ceil(args.alpha_top * N))
    set_a = set(ranked[:n_a])
    # ---- 方法B ----
    mu, sd = statistics.mean(cs.values()), statistics.stdev(cs.values())
    thr_b = mu + args.k * sd
    set_b = {n for n, v in cs.items() if v >= thr_b}
    # ---- 方法C ----
    n_cons = max(1, math.ceil(args.alpha_cons * N))
    cand = {m: set(sorted(d, key=d.get, reverse=True)[:n_cons]) for m, d in metrics.items()}
    cons = {n: sum(1 for t in cand.values() if n in t) for n in G.nodes()}
    set_c = {n for n, v in cons.items() if v >= args.cons_thr}
    # ---- 最终 D = A∩C ----
    final = set_a & set_c

    print(f"方法A(Top{args.alpha_top:.0%}): {len(set_a)} | "
          f"方法B(k={args.k}): {len(set_b)} | 方法C(共识>={args.cons_thr}): {len(set_c)} | "
          f"最终D(A∩C): {len(final)}")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # 核心基因评分表
    ws0 = wb.active
    ws0.title = "核心基因评分"
    cols0 = ["基因", "DC", "BC", "CC", "EC", "KC",
             "DĈ", "BĈ", "CĈ", "EĈ", "KĈ", "CS", "CS排名",
             "方法A", "方法B", "方法C", "最终核心(D=A∩C)", "共识指标数"]
    ws0.append(cols0)
    for n in ranked:
        ws0.append([n,
                    round(dc[n], 6), round(bc[n], 6), round(cc[n], 6), round(ec[n], 6), kc[n],
                    round(normed["DC"][n], 6), round(normed["BC"][n], 6),
                    round(normed["CC"][n], 6), round(normed["EC"][n], 6), round(normed["KC"][n], 6),
                    round(cs[n], 6), rank[n],
                    "是" if n in set_a else "", "是" if n in set_b else "",
                    "是" if n in set_c else "", "是" if n in final else "", cons[n]])
    style_header(ws0, len(cols0))
    for i, c in enumerate(cols0, 1):
        ws0.column_dimensions[get_column_letter(i)].width = 10
    ws0.column_dimensions["A"].width = 16
    ws0.freeze_panes = "A2"

    # 方法对比
    ws1 = wb.create_sheet("方法对比")
    ws1.append(["方法", "数量", "说明"])
    for k, n, desc in [("方法A: Top-N", len(set_a), f"CS排名前 max(15, {args.alpha_top:.0%}×N)"),
                       ("方法B: 统计阈值", len(set_b), f"CS ≥ μ+{args.k}σ = {thr_b:.4f}"),
                       ("方法C: 多指标共识", len(set_c), f"≥{args.cons_thr}/5 指标进入前{args.alpha_cons:.0%}"),
                       ("B∩C", len(set_b & set_c), "敏感性对照"),
                       ("最终D: A∩C", len(final), "最终核心基因")]:
        ws1.append([k, n, desc])
    style_header(ws1, 3)
    for i, c in enumerate(["方法", "数量", "说明"], 1):
        ws1.column_dimensions[get_column_letter(i)].width = 34

    # 参数
    ws2 = wb.create_sheet("参数")
    ws2.append(["参数", "值"])
    params = [
        ("输入PPI", args.input), ("权重(DC,BC,CC,EC,KC)", str(WEIGHTS)),
        ("alpha_top(方法A/最终)", args.alpha_top), ("k(方法B)", args.k),
        ("alpha_cons(方法C)", args.alpha_cons), ("cons_thr(方法C)", args.cons_thr),
        ("网络节点数(LCC)", N), ("网络边数", G.number_of_edges()),
        ("CS均值/标准差", f"{mu:.4f} / {sd:.4f}"),
        ("方法B阈值T", round(thr_b, 4)),
    ]
    for k, v in params:
        ws2.append([k, v])
    style_header(ws2, 2)
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 40

    wb.save(XLSX_OUT)
    print(f"已写出 {XLSX_OUT}")

    core_ranked = [n for n in ranked if n in final]
    with open(TXT_OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(core_ranked))
    print(f"已写出 {TXT_OUT}: {len(core_ranked)} 个核心基因")
    print("核心基因:", ", ".join(core_ranked))

    if not args.no_cytoscape:
        try:
            push_cytoscape(core_ranked, edges, PNG_OUT)
        except Exception as e:
            print("Cytoscape 不可用, 跳过子网络:", e)
    print("=== S19 核心基因筛选完成 ===")


if __name__ == "__main__":
    main()
