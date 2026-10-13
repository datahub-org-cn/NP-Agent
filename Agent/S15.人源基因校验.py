# -*- coding: utf-8 -*-
"""
S15.人源基因校验.py
===================
校验药物靶点 Gene symbol (S12) 与疾病靶点 Gene symbol (S14)
是否为【人源基因】(Homo sapiens)。

方法: 基于 NCBI 人源基因表 (Homo_sapiens.gene_info, 含 Symbol/别名/Ensembl ID)
    1) 基因名命中 NCBI 人源 Symbol 或 别名 -> 人源
    2) 以 ENSG 开头 -> 人源 Ensembl 基因ID(非标准符号, 单独标注)
    3) 以 MT- 开头 -> 人源线粒体基因(如 MT-ND6, NCBI中为 ND6)
    4) 均未命中 -> 非人源/未收录
首次运行自动下载基因表并缓存到 cashe/Homo_sapiens.gene_info.gz。

输入:
    S12.全部GeneSymbol汇总.xlsx (药物靶点, GeneSymbol汇总)
    S14.全部疾病GeneSymbol-纳入.xlsx (疾病靶点, GeneSymbol汇总)
输出:
    S15.人源基因校验.xlsx
    (工作表: 药物靶点校验 / 疾病靶点校验 / 非人源基因 / 统计)

用法: python S15.人源基因校验.py
"""
import gzip
import io as _io
import os
import re
import sys
import urllib.request
from collections import Counter

S12 = "S12.全部GeneSymbol汇总.xlsx"
S14 = "S14.全部疾病GeneSymbol-纳入.xlsx"
OUT = "S15.人源基因校验.xlsx"
GENE_INFO_URL = "https://ftp.ncbi.nlm.nih.gov/gene/DATA/GENE_INFO/Mammalia/Homo_sapiens.gene_info.gz"
GENE_INFO_CACHE = os.path.join("cashe", "Homo_sapiens.gene_info.gz")

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}


def norm_gene(v):
    s = str(v or "").strip()
    return s.upper() if s and s.upper() != "NA" else ""


def read_genes(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = {}
    for r in rows[1:]:
        d = {header[i]: ("" if v is None else v) for i, v in enumerate(r)}
        g = norm_gene(d.get("Gene symbol"))
        if not g:
            continue
        out[g] = d
    wb.close()
    return out


def ensure_gene_info():
    """下载并缓存 NCBI 人源基因表; 返回 (符号->规范名 字典, Ensembl ID集合)"""
    if not os.path.exists(GENE_INFO_CACHE):
        print("下载 NCBI 人源基因表 ...")
        req = urllib.request.Request(GENE_INFO_URL, headers=UA)
        with urllib.request.urlopen(req, timeout=180) as r:
            raw = r.read()
        os.makedirs(os.path.dirname(GENE_INFO_CACHE), exist_ok=True)
        with open(GENE_INFO_CACHE, "wb") as f:
            f.write(raw)
        print(f"已缓存 {GENE_INFO_CACHE} ({len(raw)} bytes)")
    else:
        with open(GENE_INFO_CACHE, "rb") as f:
            raw = f.read()

    symbol2canon = {}   # 小写符号/别名 -> 规范Symbol
    ensembl_ids = set()
    with gzip.open(_io.BytesIO(raw), "rt", encoding="utf-8") as f:
        for line in f:
            if line.startswith("#"):
                continue
            parts = line.rstrip("\n").split("\t")
            if len(parts) < 6:
                continue
            if parts[0] != "9606":
                continue
            symbol = parts[2].strip()
            if symbol and symbol != "-":
                symbol2canon.setdefault(symbol.lower(), symbol)
            for syn in parts[4].split("|"):
                syn = syn.strip()
                if syn and syn != "-":
                    symbol2canon.setdefault(syn.lower(), symbol or syn)
            for dbx in parts[5].split("|"):
                if dbx.startswith("Ensembl:"):
                    ensembl_ids.add(dbx.split(":", 1)[1].strip().lower())
    return symbol2canon, ensembl_ids


def validate(genes, symbol2canon, ensembl_ids):
    """返回 {gene: {human, match, note}}"""
    out = {}
    for g in genes:
        lg = g.lower()
        if lg in symbol2canon:
            out[g] = {"human": True, "match": symbol2canon[lg], "note": "NCBI人源Symbol/别名"}
        elif lg.startswith("ensg") or lg in ensembl_ids:
            out[g] = {"human": True, "match": g, "note": "人源Ensembl基因ID"}
        elif re.match(r"^mt-[a-z0-9]+$", lg):
            out[g] = {"human": True, "match": g, "note": "人源线粒体基因(MT-)"}
        else:
            out[g] = {"human": False, "match": "", "note": "未收录/非人源"}
    return out


def write_xlsx(drug_rows, dis_rows, stat):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "药物靶点校验"
    cols0 = ["Gene symbol", "是否人源", "匹配人源基因", "校验依据", "药物侧来源"]
    ws0.append(cols0)
    for r in drug_rows:
        ws0.append([r.get(c, "") for c in cols0])
    style_header(ws0, len(cols0))
    for i, c in enumerate(cols0, 1):
        ws0.column_dimensions[get_column_letter(i)].width = 16
    ws0.column_dimensions["A"].width = 18
    ws0.freeze_panes = "A2"

    ws1 = wb.create_sheet("疾病靶点校验")
    cols1 = ["Gene symbol", "是否人源", "匹配人源基因", "校验依据", "疾病侧来源"]
    ws1.append(cols1)
    for r in dis_rows:
        ws1.append([r.get(c, "") for c in cols1])
    style_header(ws1, len(cols1))
    for i, c in enumerate(cols1, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 16
    ws1.column_dimensions["A"].width = 18
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("非人源基因")
    cols2 = ["类别", "Gene symbol", "匹配人源基因", "校验依据", "来源"]
    ws2.append(cols2)
    for r in drug_rows + dis_rows:
        if r["是否人源"] != "是":
            ws2.append([r["类别"], r["Gene symbol"], r["匹配人源基因"], r["校验依据"], r["来源"]])
    style_header(ws2, len(cols2))
    for i, c in enumerate(cols2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = 16
    ws2.column_dimensions["B"].width = 18
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("统计")
    ws3.append(["统计项", "数量"])
    for k, v in stat.items():
        ws3.append([k, v])
    style_header(ws3, 2)
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 10

    wb.save(OUT)
    print(f"已写出 {OUT}")


def main():
    if not all(os.path.exists(p) for p in (S12, S14)):
        print("错误: 缺少输入文件")
        for p in (S12, S14):
            if not os.path.exists(p):
                print("  找不到:", p)
        sys.exit(1)

    drug_map = read_genes(S12, "GeneSymbol汇总")
    dis_map = read_genes(S14, "GeneSymbol汇总")
    drug_genes = list(drug_map.keys())
    dis_genes = list(dis_map.keys())
    print(f"药物靶点基因: {len(drug_genes)}; 疾病靶点基因: {len(dis_genes)}")

    symbol2canon, ensembl_ids = ensure_gene_info()
    print(f"NCBI 人源符号/别名: {len(symbol2canon)}; Ensembl ID: {len(ensembl_ids)}")

    drug_res = validate(drug_genes, symbol2canon, ensembl_ids)
    dis_res = validate(dis_genes, symbol2canon, ensembl_ids)

    drug_rows, dis_rows = [], []
    for g in drug_genes:
        r = drug_res[g]
        drug_rows.append({"Gene symbol": g, "是否人源": "是" if r["human"] else "否",
                          "匹配人源基因": r["match"], "校验依据": r["note"],
                          "来源": drug_map[g].get("来源", ""), "类别": "药物靶点"})
    for g in dis_genes:
        r = dis_res[g]
        dis_rows.append({"Gene symbol": g, "是否人源": "是" if r["human"] else "否",
                         "匹配人源基因": r["match"], "校验依据": r["note"],
                         "来源": dis_map[g].get("来源", ""), "类别": "疾病靶点"})

    n_d_yes = sum(1 for r in drug_rows if r["是否人源"] == "是")
    n_s_yes = sum(1 for r in dis_rows if r["是否人源"] == "是")
    stat = {
        "药物靶点总数": len(drug_rows), "药物靶点-人源": n_d_yes,
        "药物靶点-非人源": len(drug_rows) - n_d_yes,
        "疾病靶点总数": len(dis_rows), "疾病靶点-人源": n_s_yes,
        "疾病靶点-非人源": len(dis_rows) - n_s_yes,
    }
    print("统计:", stat)
    write_xlsx(drug_rows, dis_rows, stat)
    print("=== S15 人源基因校验完成 ===")


if __name__ == "__main__":
    main()
