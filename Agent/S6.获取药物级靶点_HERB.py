# -*- coding: utf-8 -*-
"""
S6.获取药物级靶点_HERB.py
========================
从 HERB 数据库获取复方中每味药的【药物级靶点】(herb_target)。
HERB 药物详情接口: detail_api(key_id=HERBxxxxxx, label=Herb) -> herb_target

输入: S1复方表.xlsx (中文标准名 / 中文别名, 别名用于兜底搜索)
输出: S6.药物级靶点_HERB.xlsx (工作表: 药物汇总 / 靶点明细)

用法: python S6.获取药物级靶点_HERB.py
"""
import json
import os
import re
import sys
import urllib.request
from collections import Counter

API = "http://47.92.70.12/chedi/api/"
S1_XLSX = "S1复方表.xlsx"
OUT = "S6.药物级靶点_HERB.xlsx"

DETAIL_COLS = ["药物", "HERB ID", "靶点ID", "Gene symbol", "Gene别名",
               "Protein name", "基因类型", "P值", "FDR BH", "来源"]


def call_api(func_name, **kwargs):
    data = json.dumps({"func_name": func_name, **kwargs}).encode("utf-8")
    req = urllib.request.Request(API, data=data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return json.loads(resp.read().decode("utf-8"))


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, dict):
        return str(v.get("title", ""))
    if isinstance(v, list):
        return "; ".join(to_str(x) for x in v)
    return str(v)


def search_herb(name):
    res = call_api("search_api", keyword=name, label="Herb")
    rows = res.get("res_data", [])
    for row in rows[1:]:
        if len(row) > 2 and to_str(row[2]) == name:
            return row
    if len(rows) > 1:
        return rows[1]
    return None


def find_herb(names):
    for name in names:
        herb = search_herb(name)
        if herb:
            return herb, name
    return None, ""


def read_s1(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    drugs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        cn = str(row[0]).strip()
        if not cn:
            continue
        alias_raw = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        aliases = [a for a in re.split(r"[、，,;；/／\s]+", alias_raw) if a]
        drugs.append({"标准名": cn, "别名": aliases})
    return drugs


def write_xlsx(summary_rows, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "药物汇总"
    ws0.append(["药物", "HERB ID", "匹配名称", "靶点数"])
    for r in summary_rows:
        ws0.append([r["药物"], r["HERB ID"], r["匹配名称"], r["靶点数"]])
    ws0.append(["合计", "", "", sum(r["靶点数"] for r in summary_rows)])
    style_header(ws0, 4)
    for i, c in enumerate(["药物", "HERB ID", "匹配名称", "靶点数"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14

    ws1 = wb.create_sheet("靶点明细")
    ws1.append(DETAIL_COLS)
    for r in detail_rows:
        ws1.append([r.get(c, "") for c in DETAIL_COLS])
    style_header(ws1, len(DETAIL_COLS))
    for i, c in enumerate(DETAIL_COLS, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 12
    for c in ("Gene symbol", "Gene别名", "Protein name"):
        ws1.column_dimensions[get_column_letter(DETAIL_COLS.index(c) + 1)].width = 28
    ws1.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 靶点明细 {len(detail_rows)} 行")


def main():
    if not os.path.exists(S1_XLSX):
        print(f"错误: 找不到 {S1_XLSX}")
        sys.exit(1)
    drugs = read_s1(S1_XLSX)
    print(f"读取到 {len(drugs)} 味药物")

    summary_rows, detail_rows = [], []
    for i, d in enumerate(drugs, 1):
        cn = d["标准名"]
        print(f"[{i}/{len(drugs)}] {cn} ...", flush=True)
        herb, matched = find_herb([cn] + d["别名"])
        if not herb:
            summary_rows.append({"药物": cn, "HERB ID": "", "匹配名称": "", "靶点数": 0})
            print(f"   {cn}: HERB 未找到")
            continue
        hid = to_str(herb[0])
        det = call_api("detail_api", key_id=hid, label="Herb")
        ht = det.get("herb_target", [])
        if not ht or len(ht) < 2:
            summary_rows.append({"药物": cn, "HERB ID": hid, "匹配名称": matched, "靶点数": 0})
            print(f"   {cn} -> {hid}: 无 herb_target")
            continue
        header = ht[0]
        cnt = 0
        for row in ht[1:]:
            drow = dict(zip(header, row))
            detail_rows.append({
                "药物": cn, "HERB ID": hid, "靶点ID": to_str(drow.get("Target id")),
                "Gene symbol": to_str(drow.get("Gene symbol")),
                "Gene别名": to_str(drow.get("Gene alias name")),
                "Protein name": to_str(drow.get("Protein name")),
                "基因类型": to_str(drow.get("Type of gene")),
                "P值": to_str(drow.get("P value")),
                "FDR BH": to_str(drow.get("FDR BH")),
                "来源": "HERB药物级靶点",
            })
            cnt += 1
        summary_rows.append({"药物": cn, "HERB ID": hid, "匹配名称": matched, "靶点数": cnt})
        print(f"   {cn} -> {hid}: {cnt} 个药物级靶点 (匹配: {matched})")

    write_xlsx(summary_rows, detail_rows)
    print("=== S6 HERB 药物级靶点完成 ===")


if __name__ == "__main__":
    main()
