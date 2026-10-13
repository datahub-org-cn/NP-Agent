# -*- coding: utf-8 -*-
"""
S21.核心基因-成分对应.py
========================
给出每个核心基因(S19)对应的成分列表：
即 S5 最终纳入成分中，其成分级靶点(数据库S7/S9 + 预测S11)命中该核心基因的成分。

输入:
    S19.核心基因列表.txt              (37个核心基因)
    S5.最终纳入.xlsx                  (成分, 含成分ID/PubChemCID/成分名/涉及药物)
    S4.HERB成分表-纳入.xlsx           (CID -> HERB成分ID HBIN)
    S4.TCMSP成分表-纳入.xlsx          (CID -> MOL_ID)
    S7.成分级靶点_HERB.xlsx           (HBIN -> 基因, 数据库)
    S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx (MOL_ID -> 基因, 数据库)
    S11.预测靶点_SwissTargetPrediction.xlsx  (S5成分ID -> 基因, 预测)

输出:
    S21.核心基因-成分对应.xlsx
    (工作表: 基因-成分明细 / 核心基因成分汇总 / 成分-核心基因汇总)

用法: python S21.核心基因-成分对应.py
"""
import os
import sys
from collections import defaultdict

S19 = "S19.核心基因列表.txt"
S5 = "S5.最终纳入.xlsx"
S4H = "S4.HERB成分表-纳入.xlsx"
S4T = "S4.TCMSP成分表-纳入.xlsx"
S7H = "S7.成分级靶点_HERB.xlsx"
S9T = "S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx"
S11 = "S11.预测靶点_SwissTargetPrediction.xlsx"
OUT = "S21.核心基因-成分对应.xlsx"


def read(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({hdr[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def norm(g):
    s = str(g or "").strip()
    return s.upper() if s else ""


def main():
    if not all(os.path.exists(p) for p in (S19, S5, S4H, S4T, S7H, S9T, S11)):
        print("错误: 缺少输入文件")
        for p in (S19, S5, S4H, S4T, S7H, S9T, S11):
            if not os.path.exists(p):
                print("  找不到:", p)
        sys.exit(1)

    core = []
    with open(S19, encoding="utf-8") as f:
        for line in f:
            g = line.strip()
            if g:
                core.append(g)
    core_set = {norm(g) for g in core}
    print(f"核心基因: {len(core)} 个")

    # S5 成分: CID -> 记录; 成分ID -> CID
    s5, cid5 = {}, {}
    for d in read(S5, "最终纳入"):
        cid = str(d.get("PubChem CID") or "").strip()
        i5 = str(d.get("成分ID") or "").strip()
        if cid:
            s5[cid] = d
            if i5:
                cid5[i5] = cid
    print(f"S5 成分(有CID): {len(s5)}")

    # CID -> HBIN / MOL
    cid_hbin, cid_mol = defaultdict(set), defaultdict(set)
    for d in read(S4H, "纳入成分"):
        cid = str(d.get("PubChem CID") or "").strip()
        iid = str(d.get("成分ID") or "").strip()
        if cid and iid:
            cid_hbin[cid].add(iid)
    for d in read(S4T, "纳入成分"):
        cid = str(d.get("PubChem CID") or "").strip()
        mol = str(d.get("MOL_ID") or "").strip()
        if cid and mol:
            cid_mol[cid].add(mol)

    # 成分ID -> 基因(数据库)
    hbin_genes, mol_genes = defaultdict(set), defaultdict(set)
    for d in read(S7H, "靶点明细"):
        iid = str(d.get("成分ID") or "").strip()
        g = norm(d.get("Gene symbol"))
        if iid and g:
            hbin_genes[iid].add(g)
    for d in read(S9T, "靶点明细"):
        mol = str(d.get("MOL_ID") or "").strip()
        g = norm(d.get("Gene symbol"))
        if mol and g:
            mol_genes[mol].add(g)

    # S5成分ID -> 基因(预测S11)
    stp_genes = defaultdict(set)
    from openpyxl import load_workbook
    wb11 = load_workbook(S11, read_only=True, data_only=True)
    for sheet in wb11.sheetnames:
        if sheet in ("药物汇总", "全部靶点汇总"):
            continue
        ws = wb11[sheet]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) if h is not None else "" for h in rows[0]]
        if "成分ID" not in hdr or "Gene symbol" not in hdr:
            continue
        ci, gi = hdr.index("成分ID"), hdr.index("Gene symbol")
        for r in rows[1:]:
            i5 = str(r[ci] or "").strip()
            g = norm(r[gi])
            if i5 and g and g != "N/A":
                stp_genes[i5].add(g)
    wb11.close()

    # 每个S5成分 -> {基因: 证据}
    comp_genes = {}   # cid -> {gene: "数据库"/"预测"/"数据库;预测"}
    for cid, d in s5.items():
        genes = {}
        for h in cid_hbin.get(cid, []):
            for g in hbin_genes.get(h, set()):
                genes[g] = "数据库"
        for m in cid_mol.get(cid, []):
            for g in mol_genes.get(m, set()):
                genes[g] = "数据库"
        i5 = str(d.get("成分ID") or "").strip()
        for g in stp_genes.get(i5, set()):
            genes[g] = "预测" if g not in genes else "数据库;预测"
        if genes:
            comp_genes[cid] = genes
    print(f"有靶点注释的成分: {len(comp_genes)}")

    # 核心基因 -> 成分
    gene_comps = defaultdict(list)
    for cid, genes in comp_genes.items():
        for g, ev in genes.items():
            if g in core_set:
                gene_comps[g].append((cid, ev))
    n_hit = sum(1 for g in core_set if gene_comps.get(g))
    print(f"有对应成分的核心基因: {n_hit}/{len(core)}")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # 基因-成分明细
    ws0 = wb.active
    ws0.title = "基因-成分明细"
    cols0 = ["核心基因", "成分ID", "成分名", "PubChem CID", "涉及药物", "来源库", "证据来源"]
    ws0.append(cols0)
    for g in core:
        gu = norm(g)
        comps = sorted(gene_comps.get(gu, []), key=lambda x: str(s5.get(x[0], {}).get("成分名", "")))
        for cid, ev in comps:
            d = s5.get(cid, {})
            ws0.append([g, d.get("成分ID", ""), d.get("成分名", ""), cid,
                        d.get("涉及药物", ""), d.get("来源库", ""), ev])
        if not comps:
            ws0.append([g, "", "", "", "", "", "无对应成分(无成分级/预测靶点命中)"])
    style_header(ws0, len(cols0))
    for i, c in enumerate(cols0, 1):
        ws0.column_dimensions[get_column_letter(i)].width = 16
    ws0.column_dimensions["C"].width = 40
    ws0.column_dimensions["D"].width = 40
    ws0.freeze_panes = "A2"

    # 核心基因成分汇总
    ws1 = wb.create_sheet("核心基因成分汇总")
    ws1.append(["核心基因", "对应成分数", "成分列表", "涉及药物"])
    for g in core:
        gu = norm(g)
        comps = sorted(gene_comps.get(gu, []), key=lambda x: str(s5.get(x[0], {}).get("成分名", "")))
        names = [str(s5.get(cid, {}).get("成分名", cid)) for cid, _ in comps]
        drugs = set()
        for cid, _ in comps:
            for x in str(s5.get(cid, {}).get("涉及药物") or "").split("、"):
                if x.strip():
                    drugs.add(x.strip())
        ws1.append([g, len(comps), "、".join(names), "、".join(sorted(drugs))])
    style_header(ws1, 4)
    for i, c in enumerate(["核心基因", "对应成分数", "成分列表", "涉及药物"], 1):
        ws1.column_dimensions[get_column_letter(i)].width = 16
    ws1.column_dimensions["C"].width = 60
    ws1.column_dimensions["D"].width = 40
    ws1.freeze_panes = "A2"

    # 成分-核心基因汇总
    ws2 = wb.create_sheet("成分-核心基因汇总")
    ws2.append(["成分ID", "成分名", "PubChem CID", "靶向核心基因数", "靶向核心基因", "涉及药物"])
    comp_hit = defaultdict(list)
    for g in core:
        gu = norm(g)
        for cid, _ in gene_comps.get(gu, []):
            comp_hit[cid].append(g)
    for cid in sorted(comp_hit, key=lambda c: -len(comp_hit[c])):
        d = s5.get(cid, {})
        ws2.append([d.get("成分ID", ""), d.get("成分名", ""), cid,
                    len(comp_hit[cid]), "、".join(comp_hit[cid]), d.get("涉及药物", "")])
    style_header(ws2, 6)
    for i, c in enumerate(["成分ID", "成分名", "PubChem CID", "靶向核心基因数", "靶向核心基因", "涉及药物"], 1):
        ws2.column_dimensions[get_column_letter(i)].width = 16
    ws2.column_dimensions["B"].width = 40
    ws2.column_dimensions["E"].width = 60
    ws2.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}")
    print("=== S21 核心基因-成分对应 完成 ===")


if __name__ == "__main__":
    main()
