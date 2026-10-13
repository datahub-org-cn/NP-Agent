# -*- coding: utf-8 -*-
"""
S3.筛选活性成分-DL-OB-Lipinski.py  (合并版)
=========================================
将原两步筛选合并为一步:
    原 S3.1: OB≥30% 且 DL≥0.18
    原 S3.2: Lipinski 五原则 (MW≤500, LogP≤5, HBD≤5, HBA≤10, 违规数≤1)
直接分别对 S2.HERB成分表.xlsx 与 S2.TCMSP成分表.xlsx 筛选, 各输出一份结果。

筛选规则:
    第一步 DL-OB : OB≥30% 且 DL≥0.18; OB/DL 缺失(NA/空)视为"缺参数(纳入)"
    第二步 Lipinski: 违规数≤1 为通过; MW/LogP/HBD/HBA 任一缺失视为"缺参数(纳入)"
    最终纳入 = 第一步纳入 且 第二步纳入

输入:
    S2.HERB成分表.xlsx   (工作表: 成分校验)
    S2.TCMSP成分表.xlsx  (工作表: 全部成分)

输出:
    S3.HERB成分表-OB-DL--Lipinski.xlsx   (工作表: 药物汇总 / 筛选明细 / 纳入成分)
    S3.TCMSP成分表-DL-OB-Lipinski.xlsx   (工作表: 药物汇总 / 筛选明细 / 纳入成分)

用法:
    python S3.筛选活性成分-DL-OB-Lipinski.py
"""
import os
import sys
from collections import Counter

OB_MIN = 30.0
DL_MIN = 0.18
LIPINSKI_MAX_VIOLATIONS = 1   # 允许的最大违规数

# 各数据库 S2 输入/输出及参数字段映射
DB_CONFIG = {
    "HERB": {
        "INPUT": "S2.HERB成分表.xlsx",
        "IN_SHEET": "成分校验",
        "OUT": "S3.HERB成分表-OB-DL--Lipinski.xlsx",
        "OB": "OB(%)", "DL": "DL",
        "MW": "MW", "LogP": "LogP", "HBD": "Hdon", "HBA": "Hacc",
    },
    "TCMSP": {
        "INPUT": "S2.TCMSP成分表.xlsx",
        "IN_SHEET": "全部成分",
        "OUT": "S3.TCMSP成分表-DL-OB-Lipinski.xlsx",
        "OB": "OB(%)", "DL": "DL",
        "MW": "MW", "LogP": "AlogP", "HBD": "Hdon", "HBA": "Hacc",
    },
}

EXTRA_COLS = ["DL-OB判定", "DL-OB依据", "Lipinski判定", "Lipinski违规数",
              "Lipinski违规明细", "最终结果"]


def parse_num(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.upper() == "NA":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def dl_ob_verdict(ob, dl):
    """第一步: OB≥30% 且 DL≥0.18; 缺失参数视为缺参数(纳入)"""
    ob_n, dl_n = parse_num(ob), parse_num(dl)
    if ob_n is None or dl_n is None:
        miss = []
        if ob_n is None:
            miss.append("OB缺失")
        if dl_n is None:
            miss.append("DL缺失")
        return "缺参数(纳入)", "、".join(miss)
    if ob_n >= OB_MIN and dl_n >= DL_MIN:
        return "达标", f"OB={ob_n:.2f}>=30, DL={dl_n:.3f}>=0.18"
    return "不达标", f"OB={ob_n:.2f}, DL={dl_n:.3f}"


def lipinski_verdict(r, cfg):
    """第二步: Lipinski 五原则; 任一参数缺失视为缺参数"""
    mw = parse_num(r.get(cfg["MW"]))
    logp = parse_num(r.get(cfg["LogP"]))
    hbd = parse_num(r.get(cfg["HBD"]))
    hba = parse_num(r.get(cfg["HBA"]))
    if any(v is None for v in (mw, logp, hbd, hba)):
        miss = []
        for k, v in (("MW", mw), ("LogP", logp), ("HBD", hbd), ("HBA", hba)):
            if v is None:
                miss.append(k)
        return "缺参数", -1, "缺失:" + "、".join(miss)
    viol = []
    if mw > 500:
        viol.append(f"MW={mw:.1f}>500")
    if logp > 5:
        viol.append(f"LogP={logp:.2f}>5")
    if hbd > 5:
        viol.append(f"HBD={hbd:.0f}>5")
    if hba > 10:
        viol.append(f"HBA={hba:.0f}>10")
    n = len(viol)
    verdict = "通过" if n <= LIPINSKI_MAX_VIOLATIONS else "不通过"
    detail = "; ".join(viol) if viol else "无违规"
    return verdict, n, f"{detail} (违规{n}条)"


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        d = {}
        for i, name in enumerate(header):
            if i < len(r):
                d[name] = r[i]
        out.append(d)
    return out


def write_xlsx(db_name, cfg, all_out, kept, orig_cols):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- 药物汇总 ---
    ws0 = wb.active
    ws0.title = "药物汇总"
    ws0.append(["药物", "总数", "纳入(达标)", "纳入(缺参数)", "排除"])
    drugs = []
    for r in all_out:
        if r["药物"] not in drugs:
            drugs.append(r["药物"])
    cnt_total = Counter(r["药物"] for r in all_out)
    cnt_ok = Counter(r["药物"] for r in all_out
                     if r["最终结果"] == "纳入"
                     and r["DL-OB判定"] == "达标" and r["Lipinski判定"] == "通过")
    cnt_miss = Counter(r["药物"] for r in all_out
                       if r["最终结果"] == "纳入"
                       and not (r["DL-OB判定"] == "达标" and r["Lipinski判定"] == "通过"))
    cnt_exc = Counter(r["药物"] for r in all_out if r["最终结果"] == "排除")
    for drug in drugs:
        ws0.append([drug, cnt_total.get(drug, 0), cnt_ok.get(drug, 0),
                    cnt_miss.get(drug, 0), cnt_exc.get(drug, 0)])
    n_ok = sum(cnt_ok.values())
    n_miss = len(kept) - n_ok
    ws0.append(["合计", len(all_out), n_ok, n_miss, len(all_out) - len(kept)])
    style_header(ws0, 5)
    for i, c in enumerate(["药物", "总数", "纳入(达标)", "纳入(缺参数)", "排除"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14

    # --- 筛选明细 / 纳入成分 ---
    cols = orig_cols + EXTRA_COLS
    wide_kw = ("SMILES", "InChI", "Morgan", "同义词", "别名", "IUPAC", "分子式")

    def make_sheet(title, rows):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        style_header(ws, len(cols))
        for i, c in enumerate(cols, 1):
            if c in ("InChIKey",):
                ws.column_dimensions[get_column_letter(i)].width = 34
            elif c in ("CAS",):
                ws.column_dimensions[get_column_letter(i)].width = 18
            elif any(k in c for k in wide_kw):
                ws.column_dimensions[get_column_letter(i)].width = 30
            else:
                ws.column_dimensions[get_column_letter(i)].width = 12
        ws.freeze_panes = "A2"
        return ws

    make_sheet("筛选明细", all_out)
    make_sheet("纳入成分", kept)

    wb.save(cfg["OUT"])
    print(f"已写出 {cfg['OUT']}: 纳入 {len(kept)} / {len(all_out)}")


def process(db_name, cfg):
    if not os.path.exists(cfg["INPUT"]):
        print(f"错误: 找不到 {cfg['INPUT']} (跳过 {db_name})")
        return
    rows = read_xlsx(cfg["INPUT"], cfg["IN_SHEET"])
    print(f"[{db_name}] 读取 {cfg['INPUT']} 工作表[{cfg['IN_SHEET']}]: {len(rows)} 行")
    if not rows:
        print(f"[{db_name}] 无数据")
        return
    orig_cols = list(rows[0].keys())

    all_out, kept = [], []
    for r in rows:
        ob_v, ob_d = dl_ob_verdict(r.get(cfg["OB"]), r.get(cfg["DL"]))
        lip_v, lip_n, lip_d = lipinski_verdict(r, cfg)
        keep = (ob_v in ("达标", "缺参数(纳入)")) and (lip_v in ("通过", "缺参数"))
        r["DL-OB判定"] = ob_v
        r["DL-OB依据"] = ob_d
        r["Lipinski判定"] = lip_v
        r["Lipinski违规数"] = lip_n if lip_n >= 0 else ""
        r["Lipinski违规明细"] = lip_d
        r["最终结果"] = "纳入" if keep else "排除"
        all_out.append(r)
        if keep:
            kept.append(r)

    write_xlsx(db_name, cfg, all_out, kept, orig_cols)


def main():
    for name, cfg in DB_CONFIG.items():
        process(name, cfg)
    print("=== 筛选完成 ===")


if __name__ == "__main__":
    main()
