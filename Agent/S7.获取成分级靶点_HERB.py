# -*- coding: utf-8 -*-
"""
S7.获取成分级靶点_HERB.py
=========================
从 HERB 数据库获取 S4 纳入成分的【成分级靶点】(ingredient_target)。
HERB 成分详情接口: detail_api(key_id=HBINxxxxxx, label=Ingredient) -> ingredient_target

输入: S4.HERB成分表-纳入.xlsx (工作表: 纳入成分)
输出: S7.成分级靶点_HERB.xlsx (工作表: 成分汇总 / 靶点明细)

用法: python S7.获取成分级靶点_HERB.py
"""
import json
import os
import sys
import time
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

API = "http://47.92.70.12/chedi/api/"
INPUT = "S4.HERB成分表-纳入.xlsx"
SHEET = "纳入成分"
OUT = "S7.成分级靶点_HERB.xlsx"
MAX_WORKERS = 5

DETAIL_COLS = ["药物", "成分ID", "成分名", "靶点ID", "Gene symbol", "Gene别名",
               "Protein name", "基因类型", "来源", "溯源结果"]


def call_api(func_name, **kwargs):
    data = json.dumps({"func_name": func_name, **kwargs}).encode("utf-8")
    req = urllib.request.Request(API, data=data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return json.loads(resp.read().decode("utf-8"))


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, dict):
        return str(v.get("title", ""))
    if isinstance(v, list):
        return "; ".join(to_str(x) for x in v)
    return str(v)


def get_ingredient_targets(iid):
    """按成分ID(HBIN...)获取成分级靶点列表; 失败返回 []"""
    try:
        det = call_api("detail_api", key_id=iid, label="Ingredient")
        it = det.get("ingredient_target", [])
        if not it or len(it) < 2:
            return []
        header = it[0]
        return [dict(zip(header, row)) for row in it[1:]]
    except Exception:
        return []


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def write_xlsx(summary_rows, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "成分汇总"
    ws0.append(["成分ID", "成分名", "涉及药物", "药物数", "靶点数"])
    for r in summary_rows:
        ws0.append([r["成分ID"], r["成分名"], r["涉及药物"], r["药物数"], r["靶点数"]])
    ws0.append(["合计", "", "", "", sum(r["靶点数"] for r in summary_rows)])
    style_header(ws0, 5)
    for i, c in enumerate(["成分ID", "成分名", "涉及药物", "药物数", "靶点数"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14

    ws1 = wb.create_sheet("靶点明细")
    ws1.append(DETAIL_COLS)
    for r in detail_rows:
        ws1.append([r.get(c, "") for c in DETAIL_COLS])
    style_header(ws1, len(DETAIL_COLS))
    for i, c in enumerate(DETAIL_COLS, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 12
    for c in ("Gene symbol", "Gene别名", "Protein name"):
        ws1.column_dimensions[get_column_letter(DETAIL_COLS.index(c) + 1)].width = 28
    ws1.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 靶点明细 {len(detail_rows)} 行")


def main():
    if not os.path.exists(INPUT):
        print(f"错误: 找不到 {INPUT}")
        sys.exit(1)
    rows = read_xlsx(INPUT, SHEET)
    print(f"读取 {INPUT} 工作表[{SHEET}]: {len(rows)} 行")

    # 按成分ID 去重抓取
    comp_meta = {}   # iid -> {"成分名", "药物": [...], "溯源结果"}
    for r in rows:
        iid = str(r.get("成分ID", "")).strip()
        if not iid:
            continue
        m = comp_meta.setdefault(iid, {"成分名": "", "药物": [], "溯源结果": ""})
        if not m["成分名"]:
            m["成分名"] = str(r.get("成分名", "")).strip()
        drug = str(r.get("药物", "")).strip()
        if drug and drug not in m["药物"]:
            m["药物"].append(drug)
        if not m["溯源结果"]:
            m["溯源结果"] = str(r.get("溯源结果", "")).strip()

    iids = list(comp_meta.keys())
    print(f"唯一成分ID: {len(iids)} 个, 抓取成分级靶点 ...", flush=True)
    comp_tar = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = {ex.submit(get_ingredient_targets, iid): iid for iid in iids}
        done_n = 0
        for fut in as_completed(futs):
            iid = futs[fut]
            comp_tar[iid] = fut.result()
            done_n += 1
            if done_n % 100 == 0:
                print(f"   进度 {done_n}/{len(iids)}", flush=True)

    summary_rows, detail_rows = [], []
    for iid in iids:
        meta = comp_meta[iid]
        tars = comp_tar.get(iid, [])
        for drug in meta["药物"]:
            for t in tars:
                detail_rows.append({
                    "药物": drug, "成分ID": iid, "成分名": meta["成分名"],
                    "靶点ID": to_str(t.get("Target id")),
                    "Gene symbol": to_str(t.get("Gene symbol")),
                    "Gene别名": to_str(t.get("Gene alias name")),
                    "Protein name": to_str(t.get("Protein name")),
                    "基因类型": to_str(t.get("Type of gene")),
                    "来源": to_str(t.get("Source")),
                    "溯源结果": meta["溯源结果"],
                })
        summary_rows.append({
            "成分ID": iid, "成分名": meta["成分名"],
            "涉及药物": "、".join(meta["药物"]), "药物数": len(meta["药物"]),
            "靶点数": len(tars),
        })
        if len(summary_rows) % 100 == 0:
            print(f"   汇总进度 {len(summary_rows)}/{len(iids)}", flush=True)

    write_xlsx(summary_rows, detail_rows)
    print("=== S7 HERB 成分级靶点完成 ===")


if __name__ == "__main__":
    main()
