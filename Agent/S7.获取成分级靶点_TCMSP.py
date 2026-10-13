# -*- coding: utf-8 -*-
"""
S7.获取成分级靶点_TCMSP.py
==========================
从 TCMSP 数据库获取 S4 纳入成分的【成分级靶点】。
TCMSP 分子页 molecule.php?qn=<molecule_ID> 内嵌 var tar (靶点数组)。

输入: S4.TCMSP成分表-纳入.xlsx (工作表: 纳入成分)
输出: S7.成分级靶点_TCMSP.xlsx (工作表: 成分汇总 / 靶点明细)

用法: python S7.获取成分级靶点_TCMSP.py
"""
import json
import os
import re
import sys
import time
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed

BASE = "https://www.tcmsp-e.com/"
MOLECULE_URL = BASE + "molecule.php"
INPUT = "S4.TCMSP成分表-纳入.xlsx"
SHEET = "纳入成分"
OUT = "S7.成分级靶点_TCMSP.xlsx"
MAX_WORKERS = 6

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
}

DETAIL_COLS = ["药物", "MOL_ID", "分子名", "靶点ID", "靶点名", "DrugBank ID",
               "SVM score", "RF score", "来源"]


def http_get(url, timeout=60, retries=3):
    last_err = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                raw = resp.read()
            for enc in ("utf-8", "gb18030", "latin-1"):
                try:
                    return raw.decode(enc)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", errors="replace")
        except Exception as ex:
            last_err = ex
            time.sleep(1.5 * (attempt + 1))
    raise last_err


def extract_json_array(html, key):
    m = re.search(re.escape(key) + r"\s*[=:]?\s*(\[)", html)
    if not m:
        return None
    start = m.start(1)
    depth, in_str, esc = 0, False, False
    i = start
    while i < len(html):
        ch = html[i]
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch == "[":
                depth += 1
            elif ch == "]":
                depth -= 1
                if depth == 0:
                    break
        i += 1
    try:
        return json.loads(html[start:i + 1])
    except Exception:
        return None


def get_molecule_targets(mol_id):
    num = int(mol_id[3:]) if str(mol_id).startswith("MOL") else str(mol_id)
    html = http_get(f"{MOLECULE_URL}?qn={num}")
    tar = extract_json_array(html, "var tar")
    return tar or []


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def write_xlsx(summary_rows, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "成分汇总"
    ws0.append(["MOL_ID", "分子名", "涉及药物", "药物数", "靶点数"])
    for r in summary_rows:
        ws0.append([r["MOL_ID"], r["分子名"], r["涉及药物"], r["药物数"], r["靶点数"]])
    ws0.append(["合计", "", "", "", sum(r["靶点数"] for r in summary_rows)])
    style_header(ws0, 5)
    for i, c in enumerate(["MOL_ID", "分子名", "涉及药物", "药物数", "靶点数"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14

    ws1 = wb.create_sheet("靶点明细")
    ws1.append(DETAIL_COLS)
    for r in detail_rows:
        ws1.append([r.get(c, "") for c in DETAIL_COLS])
    style_header(ws1, len(DETAIL_COLS))
    for i, c in enumerate(DETAIL_COLS, 1):
        ws1.column_dimensions[get_column_letter(i)].width = 12
    ws1.column_dimensions[get_column_letter(DETAIL_COLS.index("靶点名") + 1)].width = 34
    ws1.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 靶点明细 {len(detail_rows)} 行")


def main():
    if not os.path.exists(INPUT):
        print(f"错误: 找不到 {INPUT}")
        sys.exit(1)
    rows = read_xlsx(INPUT, SHEET)
    print(f"读取 {INPUT} 工作表[{SHEET}]: {len(rows)} 行")

    comp_meta = {}
    for r in rows:
        mol = str(r.get("MOL_ID", "")).strip()
        if not mol:
            continue
        m = comp_meta.setdefault(mol, {"分子名": "", "药物": []})
        if not m["分子名"]:
            m["分子名"] = str(r.get("分子名", "")).strip()
        drug = str(r.get("药物", "")).strip()
        if drug and drug not in m["药物"]:
            m["药物"].append(drug)

    mols = list(comp_meta.keys())
    print(f"唯一 MOL_ID: {len(mols)} 个, 抓取成分级靶点 ...", flush=True)
    mol_tar = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futs = {ex.submit(get_molecule_targets, m): m for m in mols}
        done_n = 0
        for fut in as_completed(futs):
            mol = futs[fut]
            try:
                mol_tar[mol] = fut.result()
            except Exception as e:
                mol_tar[mol] = []
                print(f"   {mol} 抓取失败: {e}", flush=True)
            done_n += 1
            if done_n % 20 == 0:
                print(f"   进度 {done_n}/{len(mols)}", flush=True)

    summary_rows, detail_rows = [], []
    for mol in mols:
        meta = comp_meta[mol]
        tars = mol_tar.get(mol, [])
        for drug in meta["药物"]:
            for t in tars:
                detail_rows.append({
                    "药物": drug, "MOL_ID": mol, "分子名": meta["分子名"],
                    "靶点ID": str(t.get("target_ID", "")),
                    "靶点名": str(t.get("target_name", "")),
                    "DrugBank ID": str(t.get("drugbank_ID", "")),
                    "SVM score": str(t.get("SVM_score", "")),
                    "RF score": str(t.get("RF_score", "")),
                    "来源": "TCMSP成分级靶点",
                })
        summary_rows.append({
            "MOL_ID": mol, "分子名": meta["分子名"],
            "涉及药物": "、".join(meta["药物"]), "药物数": len(meta["药物"]),
            "靶点数": len(tars),
        })

    write_xlsx(summary_rows, detail_rows)
    print("=== S7 TCMSP 成分级靶点完成 ===")


if __name__ == "__main__":
    main()
