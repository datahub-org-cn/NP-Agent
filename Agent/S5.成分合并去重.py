# -*- coding: utf-8 -*-
"""
S5.成分合并去重.py
==================
把 S4 去除假阳性后的两个"纳入"表合并为【一张】最终表 S5.最终纳入.xlsx，
并按化合物身份做全局去重。

输入:
    S4.HERB成分表-纳入.xlsx  (工作表: 纳入成分)
    S4.TCMSP成分表-纳入.xlsx (工作表: 纳入成分)

去重规则(任一满足即视为同一化合物, 采用并查集全局合并):
    1. PubChem CID 相同
    2. InChIKey 相同
    3. 成分名(归一化) + 分子式 相同

输出:
    S5.最终纳入.xlsx —— 单张工作表"最终纳入":
        每个唯一化合物一行, 含 PubChem CID / InChIKey / CAS / 成分名 / 分子式 /
        SMILES / InChI / IUPAC / MW / 来源库 / 涉及药物 / 药物数 / 出现次数 /
        匹配依据 / 溯源结果

用法:
    python S5.成分合并去重.py
"""
import re
from collections import Counter, defaultdict

HERB_XLSX = "S4.HERB成分表-纳入.xlsx"
TCMSP_XLSX = "S4.TCMSP成分表-纳入.xlsx"
SHEET = "纳入成分"
OUT = "S5.最终纳入.xlsx"

DRUGS = ["黄芪", "柴胡", "川芎", "水蛭", "枳实", "川牛膝", "莪术",
         "三棱", "土鳖虫", "大黄", "土茯苓", "槟榔", "赶黄草"]

OUT_COLS = ["成分ID", "PubChem CID", "InChIKey", "CAS", "成分名", "分子式",
            "Canonical SMILES", "异构SMILES", "InChI", "IUPAC名称", "MW",
            "来源库", "涉及药物", "药物数", "出现次数", "匹配依据", "溯源结果"]


class UF:
    def __init__(self, n):
        self.p = list(range(n))

    def find(self, x):
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.p[rb] = ra


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def norm_cid(v):
    s = str(v).strip()
    if not s or s.upper() == "NA":
        return ""
    m = re.search(r"\d+", s)
    return str(int(m.group(0))) if m else s


def norm_ik(v):
    return str(v).strip().upper()


def norm_name(v):
    return re.sub(r"[^a-z0-9]+", "", str(v).lower())


def norm_formula(v):
    return re.sub(r"\s+", "", str(v).strip())


def to_record(r, src):
    return {
        "src": src,
        "药物": str(r.get("药物", "")).strip(),
        "PubChem CID": norm_cid(r.get("PubChem CID", "")),
        "InChIKey": norm_ik(r.get("InChIKey", "")),
        "CAS": str(r.get("CAS", "")).strip(),
        "成分名": str(r.get("成分名") or r.get("分子名") or "").strip(),
        "分子式": norm_formula(r.get("分子式", "")),
        "Canonical SMILES": str(r.get("Canonical SMILES", "")).strip(),
        "异构SMILES": str(r.get("异构SMILES", "")).strip(),
        "InChI": str(r.get("InChI", "")).strip(),
        "IUPAC名称": str(r.get("IUPAC名称", "")).strip(),
        "MW": str(r.get("MW", "")).strip(),
        "溯源结果": str(r.get("溯源结果", "")).strip(),
    }


def merge_group(g):
    def first(key):
        return next((r[key] for r in g if r[key]), "")

    cid = first("PubChem CID")
    ik = first("InChIKey")
    cas = first("CAS")
    name = first("成分名")
    formula = first("分子式")

    drugs = []
    for r in g:
        if r["药物"] and r["药物"] not in drugs:
            drugs.append(r["药物"])
    drugs.sort(key=lambda d: DRUGS.index(d) if d in DRUGS else 999)

    srcs = sorted({r["src"] for r in g})
    src = "两库共有" if len(srcs) == 2 else srcs[0]

    # 匹配依据: 组内一致的标识符
    ev = []
    if cid and all(not r["PubChem CID"] or r["PubChem CID"] == cid for r in g):
        ev.append("PubChem CID")
    if ik and all(not r["InChIKey"] or r["InChIKey"] == ik for r in g):
        ev.append("InChIKey")
    nm, fm = norm_name(name), formula
    if nm and fm and all(
            (not norm_name(r["成分名"]) or not r["分子式"])
            or (norm_name(r["成分名"]) == nm and r["分子式"] == fm)
            for r in g):
        ev.append("名称+分子式")
    basis = "; ".join(ev) if ev else "名称归一化"

    if any(r["溯源结果"] == "溯源异常(暂保留)" for r in g):
        trace = "溯源异常(暂保留)"
    else:
        trace = "找到"

    return {
        "PubChem CID": cid, "InChIKey": ik, "CAS": cas, "成分名": name,
        "分子式": formula,
        "Canonical SMILES": first("Canonical SMILES"),
        "异构SMILES": first("异构SMILES"),
        "InChI": first("InChI"),
        "IUPAC名称": first("IUPAC名称"),
        "MW": first("MW"),
        "来源库": src, "涉及药物": "、".join(drugs), "药物数": len(drugs),
        "出现次数": len(g), "匹配依据": basis, "溯源结果": trace,
    }


def main():
    herb_rows = read_xlsx(HERB_XLSX, SHEET)
    tcmsp_rows = read_xlsx(TCMSP_XLSX, SHEET)
    print(f"读取 HERB 纳入: {len(herb_rows)} 行; TCMSP 纳入: {len(tcmsp_rows)} 行")

    records = [to_record(r, "HERB") for r in herb_rows] + \
              [to_record(r, "TCMSP") for r in tcmsp_rows]
    n = len(records)
    uf = UF(n)

    by_cid, by_ik, by_nf = defaultdict(list), defaultdict(list), defaultdict(list)
    for i, rec in enumerate(records):
        if rec["PubChem CID"]:
            by_cid[rec["PubChem CID"]].append(i)
        if rec["InChIKey"]:
            by_ik[rec["InChIKey"]].append(i)
        nm, fm = norm_name(rec["成分名"]), rec["分子式"]
        if nm and fm:
            by_nf[(nm, fm)].append(i)
    for lst in by_cid.values():
        for i in lst[1:]:
            uf.union(lst[0], i)
    for lst in by_ik.values():
        for i in lst[1:]:
            uf.union(lst[0], i)
    for lst in by_nf.values():
        for i in lst[1:]:
            uf.union(lst[0], i)

    groups = defaultdict(list)
    for i, rec in enumerate(records):
        groups[uf.find(i)].append(rec)
    print(f"合并去重后唯一化合物: {len(groups)} 个 (合并前 {n} 行)")

    out = [merge_group(g) for g in groups.values()]
    src_pri = {"两库共有": 0, "HERB": 1, "TCMSP": 2}

    def sort_key(r):
        cid = int(r["PubChem CID"]) if r["PubChem CID"].isdigit() else 10 ** 12
        return (src_pri.get(r["来源库"], 9), cid, r["成分名"])

    out.sort(key=sort_key)
    for i, r in enumerate(out, 1):
        r["成分ID"] = f"C{i:04d}"

    # 统计
    print("来源分布:", dict(Counter(r["来源库"] for r in out)))
    print("两库共有且按CID匹配示例数:",
          sum(1 for r in out if r["来源库"] == "两库共有" and "PubChem CID" in r["匹配依据"]))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "最终纳入"
    ws.append(OUT_COLS)
    for r in out:
        ws.append([r.get(c, "") for c in OUT_COLS])

    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(OUT_COLS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, c in enumerate(OUT_COLS, 1):
        if c in ("InChIKey",):
            ws.column_dimensions[get_column_letter(i)].width = 34
        elif c in ("CAS",):
            ws.column_dimensions[get_column_letter(i)].width = 18
        elif c in ("Canonical SMILES", "异构SMILES", "InChI", "IUPAC名称", "涉及药物"):
            ws.column_dimensions[get_column_letter(i)].width = 32
        else:
            ws.column_dimensions[get_column_letter(i)].width = 14
    ws.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 共 {len(out)} 个唯一化合物 (单张表: 最终纳入)")


if __name__ == "__main__":
    main()
