# -*- coding: utf-8 -*-
"""
S9.补全GeneSymbol_成分级TCMSP.py
===============================
为 S7.成分级靶点_TCMSP.xlsx 补全 Gene symbol（与 S8 同一套 TCMSP 靶点名映射）。

映射优先级:
    1. 复用 S8 已生成的缓存 cashe/s8_gene_cache.json (UniProt/mygene 映射结果)
    2. 缓存未命中时在线映射: UniProt(蛋白名相似度≥0.85) -> mygene.info(≥0.8)

输入: S7.成分级靶点_TCMSP.xlsx (工作表: 靶点明细)
输出: S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx
      (工作表: 基因映射汇总 / 靶点明细[含Gene symbol])

用法: python S9.补全GeneSymbol_成分级TCMSP.py
"""
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
from collections import Counter
from difflib import SequenceMatcher
from concurrent.futures import ThreadPoolExecutor, as_completed

INPUT = "S7.成分级靶点_TCMSP.xlsx"
SHEET = "靶点明细"
OUT = "S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx"
CACHE_FILE = "cashe/s8_gene_cache.json"   # 复用 S8 缓存
MAX_WORKERS = 5

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
UNIPROT = "https://rest.uniprot.org/uniprotkb/search"
MYGENE = "https://mygene.info/v3/query"


def _norm(s):
    return re.sub(r"[^a-z0-9]+", "", str(s).lower())


def get_json(url, timeout=40):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def uniprot_gene(name):
    q = urllib.parse.quote(f'protein_name:"{name}" AND organism_id:9606')
    d = get_json(f"{UNIPROT}?query={q}&fields=accession,gene_names,protein_name&format=json&size=10")
    target = _norm(name)
    best_score, best_gene = 0.0, ""
    for r in d.get("results", []):
        pd = r.get("proteinDescription") or {}
        rec = pd.get("recommendedName") or {}
        cands = []
        if rec.get("fullName", {}).get("value"):
            cands.append(rec["fullName"]["value"])
        if rec.get("shortName", {}).get("value"):
            cands.append(rec["shortName"]["value"])
        for alt in (pd.get("alternativeNames") or []):
            if alt.get("fullName", {}).get("value"):
                cands.append(alt["fullName"]["value"])
        score = max((SequenceMatcher(None, target, _norm(c)).ratio() for c in cands), default=0.0)
        gn = next(((g.get("geneName") or {}).get("value")
                   for g in (r.get("genes") or []) if (g.get("geneName") or {}).get("value")), "")
        if score > best_score and gn:
            best_score, best_gene = score, gn
    return best_gene if best_score >= 0.85 else ""


def mygene_symbol(name):
    q = urllib.parse.quote(name)
    d = get_json(f"{MYGENE}?q={q}&fields=symbol,name&species=human&size=5")
    target = _norm(name)
    best_score, best_sym = 0.0, ""
    for h in d.get("hits", []):
        sym = h.get("symbol")
        if not sym:
            continue
        hname = h.get("name") or ""
        score = SequenceMatcher(None, target, _norm(hname)).ratio() if hname else 0.0
        if score > best_score:
            best_score, best_sym = score, sym
    return best_sym if best_score >= 0.8 else ""


def map_gene(name):
    try:
        s = uniprot_gene(name)
        if s:
            return s, "UniProt"
    except Exception:
        pass
    try:
        s = mygene_symbol(name)
        if s:
            return s, "mygene"
    except Exception:
        pass
    return "", "未映射"


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)


def write_xlsx(map_rows, detail_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "基因映射汇总"
    ws0.append(["靶点ID", "靶点名", "Gene symbol", "映射来源"])
    for r in map_rows:
        ws0.append([r["靶点ID"], r["靶点名"], r["Gene symbol"], r["映射来源"]])
    cnt = Counter(r["映射来源"] for r in map_rows)
    ws0.append(["合计", len(map_rows),
                sum(1 for r in map_rows if r["Gene symbol"]),
                "已映射:" + str(cnt.get("UniProt", 0)) + " UniProt / " + str(cnt.get("mygene", 0)) + " mygene"])
    style_header(ws0, 4)
    for i, c in enumerate(["靶点ID", "靶点名", "Gene symbol", "映射来源"], 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14
    ws0.column_dimensions["B"].width = 46

    cols = list(detail_rows[0].keys()) if detail_rows else []
    ws1 = wb.create_sheet("靶点明细")
    ws1.append(cols)
    for r in detail_rows:
        ws1.append([r.get(c, "") for c in cols])
    style_header(ws1, len(cols))
    for i, c in enumerate(cols, 1):
        if c in ("靶点名",):
            ws1.column_dimensions[get_column_letter(i)].width = 46
        else:
            ws1.column_dimensions[get_column_letter(i)].width = 14
    ws1.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: 靶点明细 {len(detail_rows)} 行")


def main():
    if not os.path.exists(INPUT):
        print(f"错误: 找不到 {INPUT}")
        sys.exit(1)
    rows = read_xlsx(INPUT, SHEET)
    print(f"读取 {INPUT} 工作表[{SHEET}]: {len(rows)} 行")

    uniq = []
    for r in rows:
        n = str(r.get("靶点名", "")).strip()
        if n and n not in uniq:
            uniq.append(n)
    print(f"唯一靶点名: {len(uniq)} 个")

    cache = load_cache()
    todo = [n for n in uniq if n not in cache]
    print(f"缓存命中: {len(uniq) - len(todo)} 个; 需在线映射: {len(todo)} 个")

    if todo:
        done_n = 0
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
            futs = {ex.submit(map_gene, n): n for n in todo}
            for fut in as_completed(futs):
                name = futs[fut]
                try:
                    cache[name] = fut.result()
                except Exception:
                    cache[name] = ("", "异常")
                done_n += 1
                if done_n % 50 == 0:
                    print(f"   映射进度 {done_n}/{len(todo)}", flush=True)
        save_cache(cache)

    name2gene = {n: cache.get(n, ("", "未映射"))[0] for n in uniq}
    name2src = {n: cache.get(n, ("", "未映射"))[1] for n in uniq}
    for r in rows:
        n = str(r.get("靶点名", "")).strip()
        r["Gene symbol"] = name2gene.get(n, "")
        r["映射来源"] = name2src.get(n, "未映射")

    seen = {}
    for r in rows:
        tid = str(r.get("靶点ID", "")).strip()
        if tid and tid not in seen:
            seen[tid] = (str(r.get("靶点名", "")), str(r.get("Gene symbol", "")), str(r.get("映射来源", "")))
    map_rows = [{"靶点ID": tid, "靶点名": v[0], "Gene symbol": v[1], "映射来源": v[2]}
                for tid, v in seen.items()]

    cnt = Counter(r["映射来源"] for r in rows)
    print("映射来源统计(明细行):", dict(cnt))
    write_xlsx(map_rows, rows)
    print("=== S9 补全成分级靶点 Gene symbol 完成 ===")


if __name__ == "__main__":
    main()
