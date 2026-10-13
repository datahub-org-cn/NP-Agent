# -*- coding: utf-8 -*-
"""
S10.统计药物靶点数.py
====================
统计每个药物(复方中的每味药)的【药物级靶点】与【成分级靶点】数量，重复靶点按
Gene symbol 去重（无基因名的按靶点ID去重）。

数据源:
    药物级-HERB : S6.药物级靶点_HERB.xlsx                    (工作表: 靶点明细)
    药物级-TCMSP: S8.药物级靶点_TCMSP-补全GeneSymbol.xlsx    (工作表: 靶点明细)
    成分级-HERB : S7.成分级靶点_HERB.xlsx                    (工作表: 靶点明细)
    成分级-TCMSP: S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx    (工作表: 靶点明细)

输出: S10.药物靶点统计.xlsx —— 单张表"药物靶点统计"
      药物 / 药物级靶点数(去重) / 其中HERB / 其中TCMSP /
      成分级靶点数(去重) / 其中HERB / 其中TCMSP / 靶点总数(去重)

用法: python S10.统计药物靶点数.py
"""
import os
import re
import sys

S1_XLSX = "S1复方表.xlsx"
OUT = "S10.药物靶点统计.xlsx"

FILES = {
    "药物级HERB":  ("S6.药物级靶点_HERB.xlsx", "靶点明细"),
    "药物级TCMSP": ("S8.药物级靶点_TCMSP-补全GeneSymbol.xlsx", "靶点明细"),
    "成分级HERB":  ("S7.成分级靶点_HERB.xlsx", "靶点明细"),
    "成分级TCMSP": ("S9.成分级靶点_TCMSP-补全GeneSymbol.xlsx", "靶点明细"),
}

OUT_COLS = ["药物", "药物级靶点数(去重)", "药物级-HERB", "药物级-TCMSP",
            "成分级靶点数(去重)", "成分级-HERB", "成分级-TCMSP", "靶点总数(去重)"]


def read_s1(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    drugs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            cn = str(row[0]).strip()
            if cn:
                drugs.append(cn)
    return drugs


def read_xlsx(path, sheet):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({header[i]: ("" if v is None else v) for i, v in enumerate(r)})
    wb.close()
    return out


def target_key(gene, tid, prefix):
    g = str(gene or "").strip()
    if g and g.upper() != "NA":
        return ("G", g.upper())
    return ("ID", f"{prefix}:{str(tid or '').strip()}")


def collect(path, sheet, prefix):
    """返回 {药物: set(靶点key)} 与 {药物: 原始行数}"""
    rows = read_xlsx(path, sheet)
    by_drug, raw = {}, {}
    for r in rows:
        drug = str(r.get("药物", "")).strip()
        if not drug:
            continue
        key = target_key(r.get("Gene symbol"), r.get("靶点ID"), prefix)
        by_drug.setdefault(drug, set()).add(key)
        raw[drug] = raw.get(drug, 0) + 1
    return by_drug, raw


def main():
    if not os.path.exists(S1_XLSX):
        print(f"错误: 找不到 {S1_XLSX}")
        sys.exit(1)
    drugs = read_s1(S1_XLSX)
    print(f"读取到 {len(drugs)} 味药物: {'、'.join(drugs)}")

    data = {}
    for tag, (path, sheet) in FILES.items():
        if not os.path.exists(path):
            print(f"警告: 找不到 {path}, 该来源按空处理")
            data[tag] = ({}, {})
        else:
            by_drug, raw = collect(path, sheet, tag)
            data[tag] = (by_drug, raw)
            print(f"   {tag}: {sum(len(v) for v in by_drug.values())} 个去重靶点对 / {sum(raw.values())} 行")

    dh, _ = data["药物级HERB"]
    dt, _ = data["药物级TCMSP"]
    ch, _ = data["成分级HERB"]
    ct, _ = data["成分级TCMSP"]

    rows = []
    for drug in drugs:
        drug_herb = dh.get(drug, set())
        drug_tcmsp = dt.get(drug, set())
        comp_herb = ch.get(drug, set())
        comp_tcmsp = ct.get(drug, set())
        drug_level = drug_herb | drug_tcmsp
        comp_level = comp_herb | comp_tcmsp
        rows.append({
            "药物": drug,
            "药物级靶点数(去重)": len(drug_level),
            "药物级-HERB": len(drug_herb),
            "药物级-TCMSP": len(drug_tcmsp),
            "成分级靶点数(去重)": len(comp_level),
            "成分级-HERB": len(comp_herb),
            "成分级-TCMSP": len(comp_tcmsp),
            "靶点总数(去重)": len(drug_level | comp_level),
        })

    tot = [0] * 7
    for r in rows:
        for i, c in enumerate(OUT_COLS[1:], 0):
            tot[i] += r[c]
    rows.append({"药物": "合计", **{c: tot[i] for i, c in enumerate(OUT_COLS[1:])}})

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "药物靶点统计"
    ws.append(OUT_COLS)
    for r in rows:
        ws.append([r.get(c, "") for c in OUT_COLS])

    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(OUT_COLS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, c in enumerate(OUT_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "A2"

    wb.save(OUT)
    print(f"已写出 {OUT}: {len(rows)} 行 (含合计)")
    print("=== S10 统计完成 ===")


if __name__ == "__main__":
    main()
