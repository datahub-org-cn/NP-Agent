# -*- coding: utf-8 -*-
"""
S18.获取STRING-PPI网络.py
=========================
调用 STRING 数据库 (https://string-db.org/) 绘制"核心靶点"的蛋白互作网络(PPI)。

流程:
    1. 读取 S19.核心基因列表.txt 的交集基因(核心靶点)
    2. get_string_ids 校验映射, 仅保留 preferredName 与输入基因一致的条目
       (STRING 别名表会把 VEGFA->COL18A1、VDR->CYP27B1 错配; lncRNA/miRNA 无蛋白映射,
        记录到"排除基因"并剔除)
    3. 用校验后的 ENSP 以 POST 调 network 接口 (species=9606, required_score=--score)
       可选 --escore: 仅保留实验证据边(escore>0), 剔除文本挖掘/共表达等预测边
    4. 输出 xlsx: PPI边 / PPI节点 / 排除基因 / 统计
    5. 图片: 非escore版本下载 STRING 官方图; escore版本由 Cytoscape 渲染导出
       (STRING 图片接口不支持按 escore 过滤)

用法:
    python S18.获取STRING-PPI网络.py --score 900 --tag score900
    python S18.获取STRING-PPI网络.py --score 700 --tag score700
    python S18.获取STRING-PPI网络.py --score 900 --escore --tag score900_exp
    python S18.获取STRING-PPI网络.py --score 700 --escore --tag score700_exp
"""
import argparse
import json
import os
import sys
import urllib.parse
import urllib.request
from collections import Counter, defaultdict

STRING = "https://string-db.org/api"
SPECIES = 9606
CORE_TXT = "S19.核心基因列表.txt"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
CYTO = "http://127.0.0.1:1234/v1"


def rest(method, path, payload=None, timeout=180):
    url = CYTO + path
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(url, data=data, method=method,
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def http_tsv(url):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.read().decode("utf-8", "replace")


def http_tsv_post(endpoint, params):
    body = urllib.parse.urlencode(params, doseq=True).encode("utf-8")
    req = urllib.request.Request(STRING + endpoint, data=body, headers={
        "User-Agent": UA["User-Agent"],
        "Content-Type": "application/x-www-form-urlencoded",
    })
    with urllib.request.urlopen(req, timeout=300) as r:
        return r.read().decode("utf-8", "replace")


def read_core_genes_txt(path):
    with open(path, "r", encoding="utf-8") as f:
        genes = [line.strip() for line in f if line.strip()]
    return genes


def resolve_string_ids(genes):
    valid, excluded = {}, {}
    for i in range(0, len(genes), 100):
        chunk = genes[i:i + 100]
        ids = urllib.parse.quote("\r".join(chunk))
        text = http_tsv(f"{STRING}/tsv/get_string_ids?identifiers={ids}"
                        f"&species={SPECIES}&limit=5")
        lines = [l for l in text.strip().splitlines() if l.strip()]
        rows_by_q = defaultdict(list)
        for l in lines[1:]:
            p = l.split("\t")
            if len(p) >= 5:
                try:
                    rows_by_q[int(p[0])].append(p)
                except ValueError:
                    pass
        for j, sym in enumerate(chunk):
            rows = rows_by_q.get(j, [])
            matched = [p for p in rows if p[4].upper() == sym.upper()]
            if matched:
                valid[sym] = matched[0][1]
            elif rows:
                excluded[sym] = f"STRING别名冲突(被映射为{rows[0][4]})"
            else:
                excluded[sym] = "STRING无蛋白映射(非蛋白编码/未收录)"
    return valid, excluded


def fetch_network(ensps, score, escore_only=False):
    text = http_tsv_post("/tsv/network", {
        "identifiers": "\n".join(ensps), "species": SPECIES,
        "required_score": score,
    })
    lines = [l for l in text.strip().splitlines() if l.strip()]
    if not lines:
        return []
    header = lines[0].split("\t")
    edges = [dict(zip(header, l.split("\t"))) for l in lines[1:]]
    if escore_only:
        keep = []
        for e in edges:
            try:
                if float(e.get("escore") or 0) > 0:
                    keep.append(e)
            except (TypeError, ValueError):
                pass
        return keep
    return edges


def fetch_string_image(ensps, score):
    body = urllib.parse.urlencode({
        "identifiers": "\n".join(ensps), "species": SPECIES,
        "required_score": score,
    }).encode("utf-8")
    req = urllib.request.Request(STRING + "/image/network", data=body, headers={
        "User-Agent": UA["User-Agent"],
        "Content-Type": "application/x-www-form-urlencoded",
    })
    with urllib.request.urlopen(req, timeout=300) as r:
        return r.read()


def score_level(s):
    s = float(s)
    if s >= 700:
        return "高(≥700)"
    if s >= 400:
        return "中(400-699)"
    return "低(<400)"


def write_xlsx(out, edges, nodes, excluded, n_in, n_valid, isolated):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="4472C4")
    head_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    ws0 = wb.active
    ws0.title = "PPI边"
    cols = ["节点A", "节点B", "combined_score", "置信级别",
            "nscore", "fscore", "pscore", "ascore", "escore", "dscore", "tscore"]
    ws0.append(cols)
    for e in edges:
        ws0.append([e.get("preferredName_A", ""), e.get("preferredName_B", ""),
                    e.get("score", ""), score_level(e.get("score") or 0),
                    e.get("nscore", ""), e.get("fscore", ""), e.get("pscore", ""),
                    e.get("ascore", ""), e.get("escore", ""), e.get("dscore", ""),
                    e.get("tscore", "")])
    style_header(ws0, len(cols))
    for i, c in enumerate(cols, 1):
        ws0.column_dimensions[get_column_letter(i)].width = 14
    ws0.freeze_panes = "A2"

    ws1 = wb.create_sheet("PPI节点")
    ws1.append(["节点", "degree", "度分级"])
    for n, deg in sorted(nodes.items(), key=lambda x: -x[1]):
        bucket = "高(≥20)" if deg >= 20 else ("中(5-19)" if deg >= 5 else "低(<5)")
        ws1.append([n, deg, bucket])
    style_header(ws1, 3)
    for i, c in enumerate(["节点", "degree", "度分级"], 1):
        ws1.column_dimensions[get_column_letter(i)].width = 14
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("排除基因")
    ws2.append(["基因", "原因"])
    for g, reason in sorted(excluded.items()):
        ws2.append([g, reason])
    style_header(ws2, 2)
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 50

    ws2b = wb.create_sheet("游离基因(无互作)")
    ws2b.append(["基因"])
    for g in sorted(isolated):
        ws2b.append([g])
    style_header(ws2b, 1)
    ws2b.column_dimensions["A"].width = 16

    ws3 = wb.create_sheet("统计")
    ws3.append(["统计项", "数量"])
    for k, v in [("输入核心靶点数", n_in),
                 ("正确映射基因数", n_valid),
                 ("排除基因数", len(excluded)),
                 ("网络中节点数", len(nodes)),
                 ("孤立基因(无互作)", n_valid - len(nodes)),
                 ("PPI边数", len(edges)),
                 ("平均degree", round(sum(nodes.values()) / len(nodes), 2) if nodes else 0)]:
        ws3.append([k, v])
    style_header(ws3, 2)
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 10

    wb.save(out)
    print(f"已写出 {out}")


def build_cytoscape(nodes, edges, net_name, style_title, out_png=None):
    """在 Cytoscape 中建网+样式+布局, 可选导出PNG; 返回 SUID"""
    node_objs = [{"data": {"id": n, "name": n, "label": n, "degree": deg,
                           "degree_bucket": ("高" if deg >= 20 else ("中" if deg >= 5 else "低"))}}
                 for n, deg in nodes.items()]
    edge_objs = []
    for e in edges:
        s = float(e.get("score") or 0)
        sb = "高" if s >= 700 else ("中" if s >= 400 else "低")
        edge_objs.append({"data": {"source": e["preferredName_A"], "target": e["preferredName_B"],
                                   "score": s, "score_bucket": sb, "interaction": "PPI"}})
    net = json.loads(rest("POST", "/networks", {
        "data": {"name": net_name},
        "elements": {"nodes": node_objs, "edges": edge_objs},
    }).decode())
    suid = net["networkSUID"]

    style = {
        "title": style_title,
        "defaults": [
            {"visualProperty": "NODE_FILL_COLOR", "value": "#B0B0B0"},
            {"visualProperty": "NODE_LABEL_COLOR", "value": "#000000"},
            {"visualProperty": "EDGE_STROKE_UNSELECTED_PAINT", "value": "#999999"},
            {"visualProperty": "NODE_LABEL_FONT_SIZE", "value": 10},
        ],
        "mappings": [
            {"mappingType": "passthrough", "mappingColumn": "label", "mappingColumnType": "String",
             "visualProperty": "NODE_LABEL"},
            {"mappingType": "discrete", "mappingColumn": "degree_bucket", "mappingColumnType": "String",
             "visualProperty": "NODE_FILL_COLOR",
             "map": [{"key": "低", "value": "#B9D4E4"}, {"key": "中", "value": "#5B9BD5"},
                     {"key": "高", "value": "#C00000"}]},
            {"mappingType": "discrete", "mappingColumn": "degree_bucket", "mappingColumnType": "String",
             "visualProperty": "NODE_SIZE",
             "map": [{"key": "低", "value": "25"}, {"key": "中", "value": "40"}, {"key": "高", "value": "70"}]},
            {"mappingType": "discrete", "mappingColumn": "score_bucket", "mappingColumnType": "String",
             "visualProperty": "EDGE_WIDTH",
             "map": [{"key": "低", "value": "1"}, {"key": "中", "value": "3"}, {"key": "高", "value": "6"}]},
        ],
    }
    try:
        title = json.loads(rest("POST", "/styles", style).decode()).get("title", style_title)
    except Exception:
        title = style_title
    try:
        rest("POST", "/commands/vizmap/apply", {"network": f"SUID:{suid}", "styles": title})
    except Exception as e:
        print("样式应用失败(继续):", e)
    try:
        rest("POST", "/commands/layout/cose", {"network": f"SUID:{suid}"})
        views = json.loads(rest("GET", f"/networks/{suid}/views").decode())
        if views:
            rest("POST", "/commands/view/fit%20content", {"view": f"SUID:{views[0]}"})
            if out_png:
                abs_out = os.path.abspath(out_png).replace("\\", "/")
                rest("POST", "/commands/view/export%20png",
                     {"outputFile": abs_out, "view": f"SUID:{views[0]}"})
                print(f"已导出 {out_png} (Cytoscape渲染)")
    except Exception as e:
        print("布局/导出失败(继续):", e)
    return suid


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--score", type=int, default=400, help="STRING required_score (默认400)")
    ap.add_argument("--escore", action="store_true", help="仅保留实验证据边(escore>0)")
    ap.add_argument("--tag", default="", help="输出文件名后缀, 如 score900 / score900_exp")
    args = ap.parse_args()

    if not os.path.exists(CORE_TXT):
        print(f"错误: 找不到 {S16}")
        sys.exit(1)
    genes_total = read_core_genes_txt(CORE_TXT)
    print(f"核心靶点数: {len(genes_total)}")

    valid, excluded = resolve_string_ids(genes_total)
    print(f"正确映射: {len(valid)}; 排除: {len(excluded)}")
    for g, reason in list(excluded.items())[:15]:
        print(f"   排除 {g}: {reason}")
    if not valid:
        sys.exit(1)

    ensps = list(valid.values())
    edges = fetch_network(ensps, args.score, args.escore)
    tag = args.tag or f"score{args.score}{'_exp' if args.escore else ''}"
    print(f"PPI 边: {len(edges)} (score>={args.score}{', 仅实验证据' if args.escore else ''})")

    nodes = Counter()
    for e in edges:
        nodes[e.get("preferredName_A", "")] += 1
        nodes[e.get("preferredName_B", "")] += 1
    isolated = [g for g in valid if g not in nodes]
    print(f"网络中节点: {len(nodes)}; 游离节点(已去除): {len(isolated)}")

    xlsx_out = f"S18b.核心基因STRING-PPI-{tag}.xlsx"
    png_out = f"S18b.核心基因STRING-PPI-{tag}.png"
    write_xlsx(xlsx_out, edges, nodes, excluded, len(genes_total), len(valid), isolated)

    # 图片只提交"有边"的基因, 彻底去除游离节点
    connected_ensps = [valid[n] for n in nodes if n in valid]

    if args.escore:
        # STRING 图片接口不支持 escore 过滤 -> 用 Cytoscape 渲染
        try:
            build_cytoscape(nodes, edges, f"STRING-PPI-{tag}",
                            f"STRING-PPI_Style_{tag}", out_png=png_out)
        except Exception as e:
            print("Cytoscape 渲染失败:", e)
    else:
        try:
            img = fetch_string_image(connected_ensps, args.score)
            with open(png_out, "wb") as f:
                f.write(img)
            print(f"已下载 STRING 官方图 {png_out} ({len(img)} bytes)")
        except Exception as e:
            print("STRING 图片下载失败:", e)
        try:
            build_cytoscape(nodes, edges, f"STRING-PPI-{tag}", f"STRING-PPI_Style_{tag}")
        except Exception as e:
            print("Cytoscape 推送失败(跳过):", e)
    print(f"=== S18 [{tag}] 完成 ===")


if __name__ == "__main__":
    main()
